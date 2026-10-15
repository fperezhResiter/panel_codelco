"""Muestras de EDP y respaldos persistentes. La conformidad es siempre manual."""
import base64
import csv
import hashlib
import json
import re
import secrets
import sqlite3
import uuid
from collections import Counter
from contextlib import contextmanager
from datetime import date, datetime
from io import BytesIO, StringIO
from pathlib import Path
from zipfile import BadZipFile, ZipFile

from openpyxl.utils.cell import get_column_letter
from .excel import FuenteExcel, clave
from .reportes import MESES, UNIDADES, PATRON_RESIDUOS, limites_tickets, metadatos, fila_plantilla_ticket, ruta_fuente_canonica

MAX_ARCHIVO = 30 * 1024 * 1024
MAX_SOLICITUD = ((MAX_ARCHIVO + 2) // 3) * 4 + 65536
EXT_DOCUMENTOS = {'.pdf', '.png', '.jpg', '.jpeg', '.webp', '.tif', '.tiff', '.doc', '.docx', '.xls', '.xlsx', '.csv', '.txt'}


def ahora():
    return datetime.now().astimezone().isoformat(timespec='seconds')


def texto(valor):
    if valor is None:
        return ''
    if isinstance(valor, (datetime, date)):
        return valor.isoformat(sep=' ') if isinstance(valor, datetime) else valor.isoformat()
    return str(valor)


def nombre_respaldo_edp(unidad):
    return '1.1 Retiro RINSP.pdf' if unidad == 'Andina' else '2.5a TICKET INTERNO.pdf'


def nombre_respaldo_ticket(muestra, ticket):
    numero = str(ticket.get('ticket') or '').strip()
    return f'{numero}_EDP{muestra["edp"]}.pdf' if re.fullmatch(r'[0-9]+', numero) else None



def nombre_corresponde_ticket(nombre, muestra, ticket):
    esperado = nombre_respaldo_ticket(muestra, ticket)
    ruta = Path(nombre)
    partes = re.fullmatch(r'([0-9]+)[ _-]*EDP[ _-]*([0-9]+)', ruta.stem, re.IGNORECASE)
    return bool(esperado and ruta.suffix and ruta.suffix != '.' and partes
                and partes[1] == str(ticket.get('ticket') or '').strip()
                and partes[2] == str(muestra['edp']))


def orden_revision(muestra):
    # La fecha de creación no cambia al editar ni al sincronizar una revisión antigua.
    return (datetime.fromisoformat(muestra['creado']).timestamp(), muestra['id'])


def resumen_edp(fuentes, muestras):
    """Estado de la última revisión creada por EDP; las anteriores son historial."""
    grupos = {}
    vigentes = {}
    for registro in [*fuentes, *muestras]:
        llave = (registro['unidad'], registro['periodo'], registro['edp'])
        if llave not in grupos:
            grupos[llave] = {k: registro[k] for k in ('unidad', 'periodo', 'anio', 'mes', 'edp')}
            grupos[llave].update(archivos=[], revisiones=0, tickets=0, comprobados=0,
                                 muestra_vigente=None, revision_creada=None)
        grupo = grupos[llave]
        if registro['archivo'] not in grupo['archivos']:
            grupo['archivos'].append(registro['archivo'])
    for muestra in muestras:
        llave = (muestra['unidad'], muestra['periodo'], muestra['edp'])
        grupo = grupos[llave]
        grupo['revisiones'] += 1
        if llave not in vigentes or orden_revision(muestra) > orden_revision(vigentes[llave]):
            vigentes[llave] = muestra
            grupo.update(tickets=muestra['cantidad'], comprobados=muestra['comprobados'],
                         muestra_vigente=muestra['id'], revision_creada=muestra['creado'])
    for grupo in grupos.values():
        grupo['pendientes'] = grupo['tickets'] - grupo['comprobados']
        grupo['con_chequeo'] = grupo['comprobados'] > 0
        grupo['estado_codigo'] = ('sin_revision' if not grupo['con_chequeo'] else
                                  'parcial' if grupo['pendientes'] else 'completo')
        grupo['estado'] = {'sin_revision': 'Sin revisión', 'parcial': 'Chequeo parcial',
                           'completo': 'Chequeo completo'}[grupo['estado_codigo']]
        grupo['color'] = {'sin_revision': 'rojo', 'parcial': 'amarillo',
                          'completo': 'verde'}[grupo['estado_codigo']]
    return sorted(grupos.values(), key=lambda g: (g['unidad'], g['periodo'], g['edp']))


def leer_excel(contenido, unidad=None):
    # Los límites se revisan antes de descomprimir el libro.
    try:
        with ZipFile(BytesIO(contenido)) as archivo:
            if sum(x.file_size for x in archivo.infolist()) > 150 * 1024 * 1024:
                raise ValueError('El Excel descomprimido supera los 150 MB permitidos.')
    except BadZipFile as exc:
        raise ValueError('El archivo no es un Excel .xlsx o .xlsm válido.') from exc
    libro = FuenteExcel(BytesIO(contenido))
    try:
        andina = unidad == 'Andina'
        hoja = libro.hoja(r'11' if andina else PATRON_RESIDUOS)
        formulas = libro.formulas[hoja.title]
        inicio, fin, col_peso, col_ticket, tabla = limites_tickets(hoja, andina, formulas)
        col_fecha = next((c.column for c in hoja[inicio] if clave(c.value) == 'FECHA'), None)
        columnas = [c for c in range(1, hoja.max_column + 1)
                    if hoja.cell(inicio, c).value is not None or
                    any(hoja.cell(f, c).value is not None for f in range(inicio + 1, fin + 1))]
        campos = [{'columna': get_column_letter(c),
                   'nombre': texto(hoja.cell(inicio, c).value) or 'Sin encabezado'} for c in columnas]
        tickets, avisos = [], []
        if clave(hoja.cell(inicio, col_ticket).value) == 'IDMINIMIZA':
            avisos.append('Se usa ID MINIMIZA como número de ticket porque la hoja no tiene columna N.º Ticket.')
        for fila in range(inicio + 1, fin + 1):
            if andina and fila_plantilla_ticket(hoja, formulas, fila, col_ticket, col_peso, col_fecha):
                continue
            if not any(hoja.cell(fila, c).value is not None or
                       libro.formulas[hoja.title].cell(fila, c).value is not None for c in columnas):
                continue
            identificador = texto(hoja.cell(fila, col_ticket).value)
            valores = []
            for c, campo in zip(columnas, campos):
                celda = hoja.cell(fila, c)
                original = libro.formulas[hoja.title].cell(fila, c)
                valor = texto(celda.value)
                # Mantener ceros iniciales en identificadores formateados como 000000.
                if c == col_ticket and isinstance(celda.value, int) and re.fullmatch('0+', celda.number_format):
                    valor = str(celda.value).zfill(len(celda.number_format))
                    identificador = valor
                formula = original.value if original.data_type == 'f' else None
                valores.append({**campo, **({'unidad': 't'} if andina and c == col_peso else {}),
                                'valor': valor, 'celda': celda.coordinate, 'formula': formula,
                                'sin_resultado': bool(formula and celda.value is None)})
            tickets.append({'fila': fila, 'ticket': identificador, 'campos': valores})
        if not tickets:
            raise ValueError(f'La hoja {hoja.title} no contiene tickets.')
        ids = Counter(t['ticket'] for t in tickets)
        if any(n > 1 for k, n in ids.items() if k):
            avisos.append('Hay números de ticket repetidos. El sorteo distingue cada registro por su fila Excel.')
        if '' in ids:
            avisos.append('Hay filas sin número de ticket; permanecen en la población para revisión manual.')
        if any(c['sin_resultado'] for t in tickets for c in t['campos']):
            avisos.append('Hay fórmulas sin resultado guardado; se muestran como pendientes de recalcular en Excel.')
        return {'hoja': hoja.title, 'tabla': tabla, 'item': '1.1' if andina else '2.5.a',
                'tickets': tickets, 'avisos': avisos,
                'sha256': hashlib.sha256(contenido).hexdigest()}
    finally:
        libro.cerrar()


def decodificar(datos, extensiones):
    nombre = datos.get('nombre', '')
    if not isinstance(nombre, str) or not nombre.strip() or len(nombre) > 240 or any(c in nombre for c in '/\\:*?"<>|') or nombre.endswith((' ', '.')) or any(ord(c) < 32 for c in nombre):
        raise ValueError('Nombre de archivo inválido.')
    if extensiones is not None and Path(nombre).suffix.lower() not in extensiones:
        raise ValueError('Formato de archivo no admitido.')
    try:
        contenido = base64.b64decode(datos.get('contenido', ''), validate=True)
    except (ValueError, TypeError) as exc:
        raise ValueError('No se pudo leer el archivo adjunto.') from exc
    if not 0 < len(contenido) <= MAX_ARCHIVO:
        raise ValueError('Cada archivo debe contener datos y pesar como máximo 30 MB.')
    return nombre, contenido


class Auditoria:
    def __init__(self, fuentes, carpeta):
        self.fuentes = Path(fuentes)
        self.carpeta = Path(carpeta)

    @contextmanager
    def conexion(self):
        self.carpeta.mkdir(parents=True, exist_ok=True)
        db = sqlite3.connect(self.carpeta / 'auditoria.sqlite3', timeout=30)
        db.row_factory = sqlite3.Row
        try:
            db.executescript('''
                CREATE TABLE IF NOT EXISTS fuentes (id TEXT PRIMARY KEY, datos TEXT NOT NULL, contenido BLOB NOT NULL);
                CREATE TABLE IF NOT EXISTS muestras (id TEXT PRIMARY KEY, fuente TEXT NOT NULL, datos TEXT NOT NULL);
                CREATE TABLE IF NOT EXISTS documentos (id TEXT PRIMARY KEY, muestra TEXT NOT NULL,
                    fila INTEGER NOT NULL, nombre TEXT NOT NULL, creado TEXT NOT NULL, contenido BLOB NOT NULL);
            ''')
            with db:
                yield db
        finally:
            db.close()

    def catalogo(self):
        fuentes, avisos = [], []
        if self.fuentes.is_dir():
            for ruta in sorted(self.fuentes.rglob('*')):
                if not ruta.is_file() or ruta.name.startswith('~$') or ruta.suffix.lower() not in ('.xlsx', '.xlsm') or ruta.parent == self.fuentes:
                    continue
                relativa = ruta.relative_to(self.fuentes).as_posix()
                try:
                    datos = metadatos(ruta, self.fuentes)
                    fuentes.append({**datos, 'id': 'local-' + hashlib.sha256(ruta_fuente_canonica(relativa).encode()).hexdigest(),
                                    'archivo': relativa, 'nombre': ruta.name, 'origen': 'Fuentes'})
                except ValueError as exc:
                    avisos.append(f'{relativa}: {exc}')
        with self.conexion() as db:
            fuentes.extend(json.loads(r['datos']) for r in db.execute('SELECT datos FROM fuentes ORDER BY rowid'))
            muestras = [json.loads(r['datos']) for r in db.execute('SELECT datos FROM muestras ORDER BY rowid DESC')]
        resumen = [{k: m[k] for k in ('id', 'fuente', 'creado', 'unidad', 'anio', 'mes', 'periodo', 'edp', 'archivo')}
                   | {'cantidad': len(m['tickets']), 'comprobados': sum(t['comprobado'] for t in m['tickets'])} for m in muestras]
        resumen.sort(key=orden_revision, reverse=True)
        return {'fuentes': fuentes, 'muestras': resumen, 'avisos': avisos, 'unidades': list(UNIDADES),
                'estados_pago': resumen_edp(fuentes, resumen), 'conexion_bd': self.estado_conexion()}

    def reporte_unidad(self, unidad):
        if unidad not in UNIDADES:
            raise ValueError('Selecciona una unidad válida para el reporte.')
        catalogo = self.catalogo()
        with self.conexion() as db:
            # Una instantánea mantiene el resumen y el detalle coherentes ante otro check.
            db.execute('BEGIN')
            guardadas = [json.loads(r['datos']) for r in db.execute('SELECT datos FROM muestras')]
            muestras = [self.obtener(m['id'], db) for m in guardadas if m['unidad'] == unidad]
        resumen = [{**m, 'cantidad': len(m['tickets']),
                    'comprobados': sum(t['comprobado'] for t in m['tickets'])} for m in muestras]
        registros = resumen_edp([f for f in catalogo['fuentes'] if f['unidad'] == unidad], resumen)
        if not registros:
            raise ValueError('La unidad no tiene EDP ni revisiones guardadas para reportar.')
        return {'unidad': unidad, 'generado': ahora(), 'estados_pago': registros,
                'muestras': sorted(muestras, key=lambda m: (m['periodo'], m['edp'], m['creado']))}

    def importar(self, datos):
        nombre, contenido = decodificar(datos, {'.xlsx', '.xlsm'})
        unidad = datos.get('unidad')
        anio, mes = datos.get('anio'), datos.get('mes')
        edp = datos.get('edp')
        if unidad not in UNIDADES or type(anio) is not int or not 2000 <= anio <= 2100 or type(mes) is not int or not 1 <= mes <= 12:
            raise ValueError('Selecciona unidad, año y mes válidos para el Excel.')
        if type(edp) is not int or edp < 0:
            raise ValueError('Indica el número de EDP.')
        libro = leer_excel(contenido, unidad)
        identificador = 'carga-' + hashlib.sha256(f'{unidad}|{anio}|{mes}|{edp}|{libro["sha256"]}'.encode()).hexdigest()
        fuente = {'id': identificador, 'unidad': unidad, 'anio': anio, 'mes': mes, 'mes_nombre': MESES[mes],
                  'periodo': f'{anio}-{mes:02d}', 'edp': edp, 'nombre': nombre, 'archivo': nombre, 'origen': 'Cargado'}
        with self.conexion() as db:
            db.execute('INSERT OR IGNORE INTO fuentes VALUES (?,?,?)', (identificador, json.dumps(fuente), contenido))
        return fuente

    def obtener(self, identificador, db=None):
        if db is None:
            with self.conexion() as conexion:
                return self.obtener(identificador, conexion)
        fila = db.execute('SELECT datos FROM muestras WHERE id=?', (identificador,)).fetchone()
        if fila is None:
            raise ValueError('No se encontró la muestra seleccionada.')
        muestra = json.loads(fila['datos'])
        for ticket in muestra['tickets']:
            ticket['documentos'] = [dict(r) for r in db.execute(
                'SELECT id,nombre,creado,length(contenido) AS bytes FROM documentos WHERE muestra=? AND fila=? ORDER BY rowid',
                (identificador, ticket['fila']))]
        muestra['respaldo_edp'] = self.respaldo_edp(muestra, db)
        for ticket in muestra['tickets']:
            ticket['respaldo_esperado'] = nombre_respaldo_ticket(muestra, ticket)
            ticket['respaldo_ticket'] = self.respaldo_ticket(muestra, ticket, db)
        return muestra

    def carpeta_edp(self, muestra):
        if muestra.get('origen') != 'Fuentes':
            raise ValueError('El respaldo debe estar en la carpeta Fuentes del EDP correspondiente.')
        carpeta = (self.fuentes / muestra['archivo']).resolve().parent
        if not carpeta.is_relative_to(self.fuentes.resolve()):
            raise ValueError('Ruta de fuente inválida.')
        if carpeta.is_dir():
            return carpeta
        partes = Path(muestra['archivo']).parts
        if len(partes) >= 3 and partes[1] != str(muestra['anio']):
            trasladado = (self.fuentes / partes[0] / str(muestra['anio']) / Path(*partes[1:])).resolve()
            if trasladado.is_relative_to(self.fuentes.resolve()) and trasladado.is_file():
                datos = metadatos(trasladado, self.fuentes)
                if all(str(datos[k]) == str(muestra[k]) for k in ('unidad', 'anio', 'mes', 'edp')):
                    return trasladado.parent
        # Las revisiones conservan su ruta original; localizar el mismo EDP tras moverlo.
        candidatas = []
        for ruta in self.fuentes.rglob(Path(muestra['archivo']).name):
            if not ruta.is_file() or not ruta.resolve().is_relative_to(self.fuentes.resolve()):
                continue
            try:
                datos = metadatos(ruta, self.fuentes)
            except ValueError:
                continue
            if all(str(datos[k]) == str(muestra[k]) for k in ('unidad', 'anio', 'mes', 'edp')):
                candidatas.append(ruta.resolve().parent)
        if len(set(candidatas)) != 1:
            raise ValueError('No se encuentra una única carpeta Fuentes para el EDP de esta revisión.')
        return candidatas[0]

    def respaldo_ticket(self, muestra, ticket, db):
        individual = self.respaldo_individual_en_fuentes(muestra, ticket)
        if individual:
            return individual
        if ticket['comprobado']:
            guardado = next((d for d in ticket['documentos'] if d['id'] == ticket.get('respaldo_autorizado')), None)
            if guardado:
                return {**guardado, 'guardado': True, 'historico': False}
            # No revocar comprobaciones anteriores por el cambio de formato.
            anterior = next(iter(ticket['documentos']), None) or muestra.get('respaldo_edp')
            return {**anterior, 'guardado': bool(anterior.get('id')) or anterior.get('guardado', False), 'historico': True} if anterior else None
        return None

    def respaldos_individuales_en_fuentes(self, muestra, ticket):
        nombre = nombre_respaldo_ticket(muestra, ticket)
        if not nombre:
            return []
        encontrados = []
        try:
            carpeta = self.carpeta_edp(muestra)
            for candidato in sorted(carpeta.iterdir(), key=lambda p: (p.suffix.lower() != '.pdf', p.name != nombre, p.name)):
                if not nombre_corresponde_ticket(candidato.name, muestra, ticket):
                    continue
                ruta = candidato.resolve()
                if not ruta.is_relative_to(carpeta) or not ruta.is_file() or not 0 < ruta.stat().st_size <= MAX_ARCHIVO:
                    continue
                if ruta.suffix.lower() == '.pdf':
                    with ruta.open('rb') as archivo:
                        if archivo.read(5) != b'%PDF-':
                            continue
                encontrados.append({'nombre': candidato.name, 'archivo': ruta.relative_to(self.fuentes.resolve()).as_posix(),
                                    'bytes': ruta.stat().st_size, 'guardado': False, 'historico': False})
        except (ValueError, OSError):
            pass
        return encontrados

    def respaldo_individual_en_fuentes(self, muestra, ticket):
        return next(iter(self.respaldos_individuales_en_fuentes(muestra, ticket)), None)

    def leer_respaldo_ticket(self, muestra, ticket, db):
        respaldo = self.respaldo_ticket(muestra, ticket, db)
        if not respaldo:
            esperado = nombre_respaldo_ticket(muestra, ticket)
            raise ValueError(f'Falta el respaldo {esperado or "con número de ticket válido"} en la carpeta Fuentes de este EDP.')
        if respaldo.get('guardado'):
            return respaldo['nombre'], db.execute('SELECT contenido FROM documentos WHERE id=?', (respaldo['id'],)).fetchone()[0]
        ruta = (self.fuentes / respaldo['archivo']).resolve()
        if not ruta.is_relative_to(self.carpeta_edp(muestra)):
            raise ValueError('El respaldo debe pertenecer a la carpeta del EDP.')
        contenido = ruta.read_bytes()
        if not 0 < len(contenido) <= MAX_ARCHIVO or (ruta.suffix.lower() == '.pdf' and not contenido.startswith(b'%PDF-')):
            raise ValueError('El respaldo debe contener datos y pesar como máximo 30 MB; los PDF deben ser válidos.')
        return respaldo['nombre'], contenido

    def documento_ticket(self, identificador, fila):
        with self.conexion() as db:
            muestra = self.obtener(identificador, db)
            ticket = next((t for t in muestra['tickets'] if t['fila'] == fila), None)
            if not ticket:
                raise ValueError('El ticket no pertenece a esta muestra.')
            return self.leer_respaldo_ticket(muestra, ticket, db)

    def archivos_respaldo(self, reporte):
        from .auditoria_pdf import tickets_completados
        emitidos = set()
        faltantes = []
        archivos = []
        with self.conexion() as db:
            for muestra, ticket in tickets_completados(reporte['muestras']):
                respaldo = ticket.get('respaldo_ticket')
                if not respaldo:
                    faltantes.append(f'Ticket {ticket["ticket"]}, EDP {muestra["edp"]}: respaldo histórico no disponible.')
                    continue
                documentos = [self.leer_respaldo_ticket(muestra, ticket, db)]
                for individual in self.respaldos_individuales_en_fuentes(muestra, ticket):
                    if individual['nombre'] != documentos[0][0]:
                        documentos.append((individual['nombre'], (self.fuentes / individual['archivo']).read_bytes()))
                for nombre, contenido in documentos:
                    llave = (nombre, hashlib.sha256(contenido).hexdigest())
                    if llave in emitidos:
                        continue
                    emitidos.add(llave)
                    nombre = Path(nombre).name.replace('\\', '_')
                    archivos.append((f'respaldos/{muestra["id"]}/{nombre}', contenido))
        if faltantes:
            archivos.append(('respaldos-no-disponibles.txt', '\n'.join(faltantes).encode('utf-8')))
        return archivos

    def respaldo_edp(self, muestra, db):
        """Localizar el PDF común solo para conservar respaldos históricos."""
        guardado = db.execute('SELECT id,nombre,creado,length(contenido) AS bytes FROM documentos '
                              'WHERE muestra=? AND fila=0 ORDER BY rowid LIMIT 1', (muestra['id'],)).fetchone()
        if guardado:
            return {**dict(guardado), 'guardado': True}
        if muestra.get('origen') != 'Fuentes':
            return None
        try:
            carpeta = self.carpeta_edp(muestra)
        except ValueError:
            return None
        raiz = self.fuentes.resolve()
        if not carpeta.is_relative_to(raiz) or not carpeta.is_dir():
            return None
        nombres = {'11RETIRORINSP', '11RETIRORINP'} if muestra['unidad'] == 'Andina' else {'25ATICKETINTERNO'}
        for ruta in sorted(carpeta.iterdir()):
            # Las fuentes de El Salvador incluyen habitualmente "(EDP 47)".
            # Admitir únicamente el sufijo del mismo EDP, sin mezclar servicios.
            nombre = clave(ruta.stem)
            sufijo = re.search(r'(?:EDP|EP)(\d+)$', nombre)
            if sufijo and int(sufijo[1]) == muestra['edp']:
                nombre = nombre[:sufijo.start()]
            if (ruta.suffix.lower() == '.pdf' and nombre in nombres
                    and ruta.resolve().is_relative_to(carpeta) and ruta.is_file() and ruta.stat().st_size > 0):
                return {'nombre': ruta.name, 'archivo': ruta.relative_to(raiz).as_posix(),
                        'bytes': ruta.stat().st_size, 'guardado': False}
        return None

    def leer_respaldo_edp(self, muestra, db):
        respaldo = self.respaldo_edp(muestra, db)
        if not respaldo:
            nombre = nombre_respaldo_edp(muestra['unidad'])
            raise ValueError(f'No se encuentra el PDF {nombre} en la carpeta de este EDP. Actualiza la revisión o adjunta un respaldo.')
        if respaldo['guardado']:
            fila = db.execute('SELECT contenido FROM documentos WHERE id=?', (respaldo['id'],)).fetchone()
            return respaldo['nombre'], fila[0]
        ruta = (self.fuentes / respaldo['archivo']).resolve()
        carpeta = self.carpeta_edp(muestra)
        if not ruta.is_relative_to(carpeta) or not ruta.is_relative_to(self.fuentes.resolve()):
            raise ValueError('El respaldo debe pertenecer a la carpeta del EDP.')
        contenido = ruta.read_bytes()
        if not contenido:
            raise ValueError('El PDF del EDP está vacío. Adjunta un respaldo o revisa el archivo.')
        return respaldo['nombre'], contenido

    def documento_edp(self, identificador):
        with self.conexion() as db:
            return self.leer_respaldo_edp(self.obtener(identificador, db), db)

    def guardar(self, db, muestra):
        muestra['actualizado'] = ahora()
        db.execute('UPDATE muestras SET datos=? WHERE id=?', (json.dumps(muestra, ensure_ascii=False), muestra['id']))

    def publicar_cambio(self, db, muestra, fila=None):
        """Punto de extensión para compartir revisiones sin cambiar el modo local."""

    def estado_conexion(self):
        return {'modo': 'local', 'conectada': True, 'carpeta': str(self.carpeta),
                'mensaje': 'BD local conectada', 'pendientes': 0}

    def publicar_eliminacion(self, db, identificador):
        """Punto de extensión para sincronizar la eliminación de una revisión."""

    def eliminar(self, datos):
        identificador = datos.get('muestra')
        if not isinstance(identificador, str) or not identificador:
            raise ValueError('Selecciona una muestra guardada para eliminar.')
        with self.conexion() as db:
            db.execute('BEGIN IMMEDIATE')
            self.obtener(identificador, db)
            db.execute('DELETE FROM documentos WHERE muestra=?', (identificador,))
            db.execute('DELETE FROM muestras WHERE id=?', (identificador,))
            self.publicar_eliminacion(db, identificador)
        return {'eliminada': identificador}

    def leer_fuente(self, identificador):
        fuente = next((f for f in self.catalogo()['fuentes'] if f['id'] == identificador), None)
        if fuente is None:
            raise ValueError('Selecciona un Excel EDP disponible.')
        if fuente['origen'] == 'Fuentes':
            ruta = (self.fuentes / fuente['archivo']).resolve()
            if not ruta.is_relative_to(self.fuentes.resolve()):
                raise ValueError('Ruta de fuente inválida.')
            if ruta.stat().st_size > MAX_ARCHIVO:
                raise ValueError('El Excel supera los 30 MB permitidos.')
            contenido = ruta.read_bytes()
        else:
            with self.conexion() as db:
                contenido = db.execute('SELECT contenido FROM fuentes WHERE id=?', (fuente['id'],)).fetchone()[0]
        return fuente, leer_excel(contenido, fuente['unidad'])

    def listar_tickets(self, identificador):
        fuente, libro = self.leer_fuente(identificador)
        return {'fuente': fuente['id'], 'sha256': libro['sha256'], 'hoja': libro['hoja'], 'item': libro['item'],
                'tickets': [{'fila': t['fila'], 'ticket': t['ticket'],
                             'fecha': next((c['valor'] for c in t['campos'] if clave(c['nombre']) == 'FECHA'), '')}
                            for t in libro['tickets']]}

    def seleccionar(self, datos, guardar=True):
        filas = datos.get('filas')
        if 'filas' in datos:
            if (not isinstance(filas, list) or not 3 <= len(filas) <= 5 or
                    any(type(fila) is not int for fila in filas) or len(set(filas)) != len(filas)):
                raise ValueError('Selecciona entre 3 y 5 tickets distintos de la lista.')
        else:
            # Compatibilidad con la selección individual de versiones anteriores.
            if type(datos.get('fila')) is not int:
                raise ValueError('Selecciona un ticket de la lista.')
            filas = [datos['fila']]
        fuente, libro = self.leer_fuente(datos.get('fuente'))
        if datos.get('sha256') != libro['sha256']:
            raise ValueError('El Excel cambió desde que se listaron los tickets. Vuelve a cargar la lista.')
        por_fila = {t['fila']: t for t in libro['tickets']}
        if any(fila not in por_fila for fila in filas):
            raise ValueError('El ticket seleccionado no pertenece a este EDP.')
        return self.crear_muestra(fuente, libro, [por_fila[fila] for fila in filas], 'manual', guardar)

    def sortear(self, datos, guardar=True):
        cantidad = datos.get('cantidad')
        if type(cantidad) is not int or cantidad not in (3, 4, 5):
            raise ValueError('Elige una muestra de 3, 4 o 5 tickets.')
        fuente, libro = self.leer_fuente(datos.get('fuente'))
        if len(libro['tickets']) < cantidad:
            raise ValueError(f'Este EDP tiene {len(libro["tickets"])} registros; no alcanza para sortear {cantidad} sin repetir filas.')
        tickets = secrets.SystemRandom().sample(libro['tickets'], cantidad)
        return self.crear_muestra(fuente, libro, tickets, 'aleatoria', guardar)

    def crear_muestra(self, fuente, libro, tickets, modo, guardar=True):
        for t in tickets:
            t.update(comprobado=False, comprobado_en=None, documentos=[])
        muestra = {**fuente, 'id': uuid.uuid4().hex, 'fuente': fuente['id'],
                   'creado': datetime.now().astimezone().isoformat(timespec='microseconds'), 'actualizado': ahora(),
                   'hoja': libro['hoja'], 'item': libro['item'], 'sha256': libro['sha256'], 'poblacion': len(libro['tickets']),
                   'avisos': libro['avisos'], 'tickets': tickets, 'modo': modo}
        with self.conexion() as db:
            if not guardar:
                muestra['respaldo_edp'] = self.respaldo_edp(muestra, db)
                muestra['preview'] = True
                return muestra
            db.execute('INSERT INTO muestras VALUES (?,?,?)', (muestra['id'], fuente['id'], json.dumps(muestra, ensure_ascii=False)))
            self.publicar_cambio(db, self.obtener(muestra['id'], db))
        return self.obtener(muestra['id'])

    def modificar(self, datos, accion):
        identificador = datos.get('muestra')
        fila = datos.get('fila')
        if type(fila) is not int:
            raise ValueError('Fila de ticket inválida.')
        archivo = decodificar(datos, None) if accion == 'adjuntar' else None
        with self.conexion() as db:
            # Serializar cambios concurrentes para no perder casillas ni respaldos.
            db.execute('BEGIN IMMEDIATE')
            muestra = self.obtener(identificador, db)
            ticket = next((t for t in muestra['tickets'] if t['fila'] == fila), None)
            if ticket is None:
                raise ValueError('El ticket no pertenece a esta muestra.')
            if accion == 'comprobar':
                if type(datos.get('comprobado')) is not bool:
                    raise ValueError('La comprobación debe ser una casilla válida.')
                if datos['comprobado'] and not ticket['comprobado']:
                    nombre, contenido = self.leer_respaldo_ticket(muestra, ticket, db)
                    existente = db.execute('SELECT id FROM documentos WHERE muestra=? AND fila=? AND nombre=? AND contenido=? LIMIT 1',
                                           (identificador, fila, nombre, contenido)).fetchone()
                    documento = existente['id'] if existente else uuid.uuid4().hex
                    if not existente:
                        db.execute('INSERT INTO documentos VALUES (?,?,?,?,?,?)',
                                   (documento, identificador, fila, nombre, ahora(), contenido))
                    ticket['respaldo_autorizado'] = documento
                if datos['comprobado'] and ticket['comprobado']:
                    return muestra
                ticket['comprobado'] = datos['comprobado']
                ticket['comprobado_en'] = ahora() if ticket['comprobado'] else None
            elif accion == 'adjuntar':
                nombre, contenido = archivo
                esperado = nombre_respaldo_ticket(muestra, ticket)
                if not nombre_corresponde_ticket(nombre, muestra, ticket):
                    raise ValueError(f'Sube un archivo cuyo nombre base sea exactamente {Path(esperado).stem if esperado else "NumeroTicket_EDPNumeroEDP"} o la variante NumeroTicket EDP NumeroEDP, con o sin guion, y la extensión de su formato.')
                if Path(nombre).suffix.lower() == '.pdf' and not contenido.startswith(b'%PDF-'):
                    raise ValueError('El archivo PDF debe ser válido.')
                if ticket['comprobado']:
                    raise ValueError('Desmarca la comprobación antes de cambiar el respaldo de este ticket.')
                carpeta = self.carpeta_edp(muestra)
                ruta = (carpeta / nombre).resolve()
                if not ruta.is_relative_to(carpeta):
                    raise ValueError('Ruta de respaldo inválida.')
                if ruta.exists():
                    if ruta.read_bytes() != contenido:
                        raise ValueError('Ya existe un respaldo con este nombre y contenido diferente en Fuentes. No se ha reemplazado.')
                else:
                    with ruta.open('xb') as destino:
                        destino.write(contenido)
                db.execute('INSERT INTO documentos VALUES (?,?,?,?,?,?)',
                           (uuid.uuid4().hex, identificador, fila, nombre, ahora(), contenido))
                ticket.update(comprobado=False, comprobado_en=None)
            elif accion == 'quitar':
                cursor = db.execute('DELETE FROM documentos WHERE id=? AND muestra=? AND fila=?', (datos.get('documento'), identificador, fila))
                if cursor.rowcount != 1:
                    raise ValueError('No se encontró el documento en este ticket.')
                ticket.update(comprobado=False, comprobado_en=None)
            else:
                raise ValueError('Acción no reconocida.')
            self.guardar(db, muestra)
            resultado = self.obtener(identificador, db)
            self.publicar_cambio(db, resultado, fila)
            return resultado

    def documento(self, identificador):
        with self.conexion() as db:
            fila = db.execute('SELECT nombre,contenido FROM documentos WHERE id=?', (identificador,)).fetchone()
            if fila is None:
                raise ValueError('Documento no encontrado.')
            return fila['nombre'], fila['contenido']


def crear_csv(muestra):
    salida = StringIO(newline='')
    escritor = csv.writer(salida, delimiter=';', lineterminator='\r\n')
    def segura(v):
        v = texto(v)
        return "'" + v if v.lstrip().startswith(('=', '+', '-', '@')) or v.startswith(('\t', '\r', '\n')) else v
    columnas = [(c['columna'], c['nombre']) for c in muestra['tickets'][0]['campos']]
    escritor.writerow(['Unidad', 'Año', 'Mes', 'Período', 'EDP', 'Excel', 'Hoja', 'Muestra', 'Creada el',
                       'Población', 'Tamaño muestra', 'SHA256 Excel', 'Fila Excel', 'Ticket'] +
                      [segura(f'{n}' + (' (t)' if clave(n) == 'CANTIDAD' and muestra['unidad'] == 'Andina' else '') + f' [{c}]') for c, n in columnas] +
                      ['Estado revisión manual', 'Comprobado el', 'Documentos', 'Observaciones de lectura', 'Tipo de selección', 'Respaldo del EDP', 'Ítem'])
    for t in muestra['tickets']:
        campos = {c['columna']: c for c in t['campos']}
        valores = [campos[c]['valor'] if not campos[c]['sin_resultado'] else 'Sin resultado guardado' for c, _ in columnas]
        escritor.writerow([segura(v) for v in [muestra['unidad'], muestra['anio'], muestra['mes'], muestra['periodo'],
            muestra['edp'], muestra['archivo'], muestra['hoja'], muestra['id'], muestra['creado'], muestra['poblacion'],
            len(muestra['tickets']), muestra['sha256'], t['fila'], t['ticket'], *valores,
            'Comprobado manualmente' if t['comprobado'] else 'Pendiente de comprobación', t['comprobado_en'],
            ' | '.join(d['nombre'] for d in t['documentos']), ' | '.join(muestra['avisos']),
            muestra.get('modo', 'aleatoria'), (t.get('respaldo_ticket') or muestra.get('respaldo_edp') or {}).get('nombre', ''),
            muestra.get('item') or ('1.1' if clave(muestra['hoja']) == '11' else '2.5.a')]])
    return ('\ufeff' + salida.getvalue()).encode('utf-8')
