"""Conciliación de Andina (1.1) y El Salvador (2.5.a) frente a sus tickets."""
import re
from copy import deepcopy
from collections import Counter, defaultdict
from datetime import datetime
from decimal import Decimal, ROUND_HALF_UP
from pathlib import Path
from openpyxl.utils.cell import range_boundaries, get_column_letter
from .excel import FuenteExcel, clave, normalizar, fecha, unico, columna_precio, columna_edp
from .andina import leer_avance_andina

MESES = ('', 'Enero', 'Febrero', 'Marzo', 'Abril', 'Mayo', 'Junio',
         'Julio', 'Agosto', 'Septiembre', 'Octubre', 'Noviembre', 'Diciembre')
UNIDADES = ('Andina', 'El Salvador')
PATRON_DETALLE = r'DETALLEDEIMPUTACIONES(?:(?:EDP|EP)?N?O?\d+)?'
PATRON_RESIDUOS = r'RESN OPELIGROSOS?'.replace(' ', '')
ERROR_EDP_DUPLICADO = 'Hay más de un Excel para la misma unidad, período y EDP. Deja una única versión para conciliar.'
VERSIONES_LECTURA = {'Andina': 'andina-1.1-v2', 'El Salvador': 'salvador-2.5.a-v4'}


def metadatos(ruta, raiz):
    partes = ruta.relative_to(raiz).parts
    if len(partes) < 3:
        raise ValueError('Ubica el Excel dentro de Unidad / Año-Mes EDP N.º.')
    unidad = next((u for u in UNIDADES if clave(u) in clave(partes[0])), None)
    if not unidad:
        raise ValueError('La carpeta de unidad debe identificar Andina o El Salvador.')
    carpetas = ' / '.join(partes[1:-1])
    periodo = re.search(r'(?<!\d)(20\d{2})\s*[-_/ ]\s*(0?[1-9]|1[0-2])(?!\d)', carpetas)
    if periodo:
        anio, mes = map(int, periodo.groups())
    else:
        anios = re.findall(r'(?<!\d)(20\d{2})(?!\d)', carpetas)
        meses = [i for i, m in enumerate(MESES) if m and normalizar(m) in normalizar(carpetas)]
        if len(set(anios)) != 1 or len(meses) != 1:
            raise ValueError('No se reconoce el año y mes de la carpeta; usa 2026 -05 EDP 47.')
        anio, mes = int(anios[0]), meses[0]
    edps = re.findall(r'(?:EDP|EP)\s*(?:N[°ºO.]?\s*)?(\d+)', normalizar(carpetas))
    edp = int(unico(list(set(edps)), 'No se reconoce un único número de EDP en la carpeta.'))
    return {'unidad': unidad, 'anio': anio, 'mes': mes, 'mes_nombre': MESES[mes],
            'periodo': f'{anio}-{mes:02d}', 'edp': edp}


def leer_imputacion(libro, edp):
    hoja = libro.hoja(PATRON_DETALLE)
    celdas = [c for fila in hoja for c in fila if c.value is not None]
    item = unico([c for c in celdas if re.fullmatch(r'2\.5\.?A\.?', normalizar(c.value))],
                 'No se encuentra un único ítem 2.5.a en DETALLE DE IMPUTACIONES.')
    cabeceras = [c for c in celdas if c.row < item.row]
    mod = unico([c for c in cabeceras if re.fullmatch(r'MODIFICACIONN?O?1', clave(c.value))],
                'No se identifica Modificación N.º 1.')
    fisico = unico([c for c in cabeceras if clave(c.value).startswith('AVANCEFISICO')],
                   'No se identifica el bloque Avance físico.')
    financiero = unico([c for c in cabeceras if clave(c.value) == 'AVANCEFINANCIERO'],
                        'No se identifica el bloque Avance financiero.')
    columnas = {'precio': columna_precio(hoja, mod, item.row),
                'ton_edp': columna_edp(hoja, fisico, item.row, edp),
                'monto_edp': columna_edp(hoja, financiero, item.row, edp)}
    valores = {k: libro.decimal(hoja, item.row, col) for k, col in columnas.items()}
    refs = {k: libro.referencia(hoja, item.row, col) for k, col in columnas.items()}
    periodo_pago = ''
    for c in cabeceras:
        if 'PERIODODEPAGO' in clave(c.value):
            for col in range(c.column + 1, hoja.max_column + 1):
                v = hoja.cell(c.row, col).value
                if v is not None:
                    periodo_pago = str(v).strip()
                    break
    return valores, refs, periodo_pago


def limites_tickets(hoja, en_toneladas=False, formulas=None):
    candidatos = []
    candidatos_minimiza = []
    nombres_peso = ('CANTIDAD', 'CANTIDADTON', 'CANTIDADTONELADAS') if en_toneladas else ('PESOTOTAL', 'PESOTOTALKG', 'PESOTOTALKGS')
    for fila in hoja.iter_rows(max_row=min(hoja.max_row, 60)):
        pesos = [c for c in fila if clave(c.value) in nombres_peso]
        tickets = [c for c in fila if 'TICKET' in clave(c.value)]
        if len(pesos) == 1 and len(tickets) == 1:
            candidatos.append((fila[0].row, pesos[0].column, tickets[0].column))
        elif not en_toneladas and len(pesos) == 1 and not tickets:
            minimiza = [c for c in fila if clave(c.value) == 'IDMINIMIZA']
            if len(minimiza) == 1:
                candidatos_minimiza.append((fila[0].row, pesos[0].column, minimiza[0].column))
    if not candidatos:
        candidatos = candidatos_minimiza
    cabecera, peso, ticket = unico(candidatos, 'No se identifica una tabla única con Ticket y ' +
                                 ('Cantidad (toneladas).' if en_toneladas else 'Peso Total (kg).'))
    tablas = []
    for tabla in hoja.tables.values():
        izquierda, arriba, derecha, abajo = range_boundaries(tabla.ref)
        if arriba == cabecera and izquierda <= peso <= derecha and izquierda <= ticket <= derecha:
            tablas.append((abajo - (tabla.totalsRowCount or 0), tabla.name))
    if tablas:
        fin, nombre = unico(tablas, 'Hay varias tablas de tickets superpuestas.')
    else:
        # Fallback explícito: terminar en el primer total, nunca sumar resúmenes inferiores.
        fin, nombre = None, None
    # Algunas tablas contienen la fila TOTAL sin declararla como totalsRowCount.
    # Los formatos antiguos de Andina usan una SUM sin etiqueta al pie.
    for fila in hoja.iter_rows(min_row=cabecera + 1, max_row=fin or hoja.max_row):
        formula = formulas.cell(fila[0].row, peso).value if formulas is not None else None
        letra = get_column_letter(peso)
        rango_total = rf'\$?{letra}\$?{cabecera + 1}:\$?{letra}\$?{fila[0].row - 1}'
        suma_final = (en_toneladas and hoja.cell(fila[0].row, ticket).value is None and
                      isinstance(formula, str) and re.fullmatch(
                          rf'=\+?(?:SUM\(|SUBTOTAL\(\d+[,;]){rango_total}\)', formula.replace(' ', ''), re.I))
        if suma_final or any(re.match(r'^(?:SUB)?TOTAL(?:KG|TON|ES)?$', clave(c.value)) for c in fila):
            fin = fila[0].row - 1
            break
    if fin is None:
        raise ValueError('La hoja no tiene tabla Excel ni fila TOTAL que delimite los tickets.')
    return cabecera, fin, peso, ticket, nombre


def fila_plantilla_ticket(hoja, formulas, fila, ticket_col, peso_col, fecha_col):
    """Filas precargadas de Andina que aún no contienen un registro de ticket."""
    return all(hoja.cell(fila, col).value is None and formulas.cell(fila, col).data_type != 'f'
               for col in (ticket_col, peso_col, fecha_col or ticket_col))


def leer_tickets(libro, unidad=None):
    en_toneladas = unidad == 'Andina'
    hoja = libro.hoja(r'11' if en_toneladas else PATRON_RESIDUOS)
    inicio, fin, peso_col, ticket_col, tabla = limites_tickets(hoja, en_toneladas, libro.formulas[hoja.title])
    encabezados = {clave(c.value): c.column for c in hoja[inicio] if c.value is not None}
    fecha_col = encabezados.get('FECHA')
    retiro_col = next((col for k, col in encabezados.items() if k == 'LUGARDERETIRO' or k.startswith('PUNTODERETIRO')), None)
    residuos_col = next((col for k, col in encabezados.items() if k.startswith(('TIPODERESIDUO', 'TIPORESIDUO'))), None)
    if residuos_col is None and en_toneladas:
        residuos_col = next((col for k, col in encabezados.items() if k.startswith(
            ('TIPODEPRODUCTO', 'TIPOPRODUCTO')) or k == 'PRODUCTO'), None)
    total = Decimal(0)
    tickets, errores, avisos, ids = [], [], [], []
    usa_minimiza = clave(hoja.cell(inicio, ticket_col).value) == 'IDMINIMIZA'
    if usa_minimiza:
        avisos.append('Se usa ID MINIMIZA como número de ticket porque la hoja no tiene columna N.º Ticket.')
    for fila in range(inicio + 1, fin + 1):
        if not any(c.value is not None for c in hoja[fila]):
            continue
        if en_toneladas and fila_plantilla_ticket(hoja, libro.formulas[hoja.title], fila,
                                                ticket_col, peso_col, fecha_col):
            # Plantillas con precio y Nr. EDP precargados, sin un ticket registrado.
            continue
        id_ticket = hoja.cell(fila, ticket_col).value
        if id_ticket is None or str(id_ticket).strip() == '':
            errores.append(f'Fila {fila}: falta ID MINIMIZA (usado como número de ticket).' if usa_minimiza
                           else f'Fila {fila}: falta el número de ticket.')
        else:
            ids.append(str(id_ticket).strip())
        try:
            cantidad = libro.decimal(hoja, fila, peso_col)
            peso = cantidad * Decimal(1000) if en_toneladas else cantidad
            total += peso
        except ValueError as exc:
            peso = None
            errores.append(str(exc))
        fecha_valor = hoja.cell(fila, fecha_col).value if fecha_col else None
        tickets.append({'fila': fila, 'ticket': str(id_ticket) if id_ticket is not None else '',
                        'fecha': fecha(fecha_valor), 'peso_kg': float(peso) if peso is not None else None,
                        'retiro': str(hoja.cell(fila, retiro_col).value or '') if retiro_col else '',
                        'residuo': str(hoja.cell(fila, residuos_col).value or '') if residuos_col else '',
                        'celda': f'{get_column_letter(peso_col)}{fila}',
                        **({'peso_ton': float(cantidad) if peso is not None else None} if en_toneladas else {})})
    if not tickets:
        errores.append('La tabla de tickets está vacía; no se interpreta como cero toneladas.')
    repetidos = [n for n, cantidad in Counter(ids).items() if cantidad > 1]
    if repetidos:
        avisos.append('Tickets repetidos: ' + ', '.join(repetidos) + '. Se incluyen todas las filas; no se eliminan duplicados.')
    if '0' in ids:
        avisos.append('Hay tickets con número 0; su peso está incluido.')
    if any(t['fecha'] is None for t in tickets):
        avisos.append('Hay tickets sin fecha reconocible; su peso está incluido.')
    if any(t['peso_kg'] is not None and t['peso_kg'] < 0 for t in tickets):
        avisos.append('Hay pesos negativos; se incluyen con su signo.')
    referencia = {'hoja': hoja.title, 'rango': f'{get_column_letter(peso_col)}{inicio + 1}:{get_column_letter(peso_col)}{fin}',
                  'tabla': tabla, 'unidad': 't' if en_toneladas else 'kg'}
    if usa_minimiza:
        referencia.update(identificador='ID MINIMIZA', columna_identificador=get_column_letter(ticket_col))
    return total if not errores else None, tickets, referencia, errores, avisos


def comparar(precio, ton_edp, monto_edp, peso_kg):
    toneladas = peso_kg / Decimal(1000)
    esperado = (precio * toneladas).quantize(Decimal('0.01'), rounding=ROUND_HALF_UP)
    monto = monto_edp.quantize(Decimal('0.01'), rounding=ROUND_HALF_UP)
    diferencia = monto - esperado
    delta_ton = ton_edp - toneladas
    # Tolerancias exclusivamente técnicas para valores binarios guardados por Excel.
    cuadra = diferencia == 0 and abs(delta_ton) <= Decimal('0.000001')
    return {'peso_kg': float(peso_kg), 'ton_tickets': float(toneladas),
            'monto_esperado': float(esperado), 'diferencia': float(diferencia),
            'diferencia_ton': float(delta_ton), 'estado': 'cuadra' if cuadra else 'diferencia'}


def crear_reporte(raiz, reutilizables=None, firmas=None):
    raiz = Path(raiz)
    reutilizables = reutilizables or {}
    firmas = firmas or {}
    registros, avisos, omitidos = [], [], []
    if not raiz.is_dir():
        avisos.append('No existe la carpeta Fuentes. Configura --fuentes o CARPETA_FUENTES.')
    else:
        for ruta in sorted(raiz.rglob('*')):
            if not ruta.is_file() or ruta.name.startswith('~$') or ruta.suffix.lower() not in ('.xlsx', '.xlsm', '.xls'):
                continue
            relativa = ruta.relative_to(raiz).as_posix()
            # Los informes generales en la raíz no representan estados de pago.
            if ruta.parent == raiz:
                omitidos.append(relativa)
                continue
            guardado = reutilizables.get(relativa)
            firma_actual = firmas.get(relativa)
            if (guardado and firma_actual is not None and guardado[0] == firma_actual and
                    guardado[1].get('version_lectura') == VERSIONES_LECTURA.get(guardado[1].get('unidad'))):
                registro = deepcopy(guardado[1])
                registro.setdefault('item', '1.1' if registro.get('unidad') == 'Andina' else '2.5.a')
                if not any('No se pudo leer el archivo (' in str(e) for e in registro.get('errores', [])):
                    registro['errores'] = [e for e in registro.get('errores', []) if e != ERROR_EDP_DUPLICADO]
                    if registro.get('monto_esperado') is not None and registro.get('diferencia_ton') is not None:
                        registro['estado'] = ('cuadra' if registro.get('diferencia') == 0 and
                                              abs(Decimal(str(registro['diferencia_ton']))) <= Decimal('0.000001')
                                              else 'diferencia')
                    else:
                        registro['estado'] = 'incompleto'
                    registros.append(registro)
                    continue
            registro = {'archivo': relativa, 'nombre': ruta.name, 'unidad': None, 'anio': None,
                        'mes': None, 'periodo': None, 'edp': None, 'estado': 'incompleto',
                        'errores': [], 'avisos': [], 'tickets': [], 'referencias': {},
                        'precio': None, 'ton_edp': None, 'monto_edp': None, 'peso_kg': None,
                        'ton_tickets': None, 'monto_esperado': None, 'diferencia': None,
                        'diferencia_ton': None, 'periodo_pago': ''}
            libro = None
            try:
                registro.update(metadatos(ruta, raiz))
                registro['item'] = '1.1' if registro['unidad'] == 'Andina' else '2.5.a'
                registro['version_lectura'] = VERSIONES_LECTURA[registro['unidad']]
                if ruta.suffix.lower() == '.xls':
                    raise ValueError('Formato .xls no compatible. Guarda una copia en .xlsx.')
                libro = FuenteExcel(ruta)
                # Guardar tickets aun cuando falle la imputación: densidades depende de
                # peso y residuo, no de la conciliación financiera del EDP.
                peso, tickets, ref, errores, advertencias = leer_tickets(libro, registro['unidad'])
                registro['tickets'] = tickets
                registro['referencias']['peso'] = ref
                registro['errores'].extend(errores)
                registro['avisos'].extend(advertencias)
                valores, refs, periodo_pago = (leer_avance_andina(libro) if registro['unidad'] == 'Andina'
                                               else leer_imputacion(libro, registro['edp']))
                registro.update({k: float(v) for k, v in valores.items()})
                registro['referencias'].update(refs)
                registro['periodo_pago'] = periodo_pago
                if peso is not None:
                    registro.update(comparar(valores['precio'], valores['ton_edp'], valores['monto_edp'], peso))
            except Exception as exc:
                registro['errores'].append(str(exc) if isinstance(exc, ValueError) else
                                           f'No se pudo leer el archivo ({type(exc).__name__}). Comprueba que está descargado y accesible.')
            finally:
                if libro:
                    libro.cerrar()
            registros.append(registro)
    grupos = defaultdict(list)
    for r in registros:
        if r['unidad'] and r['periodo'] and r['edp'] is not None:
            grupos[(r['unidad'], r['periodo'], r['edp'])].append(r)
    for grupo in grupos.values():
        if len(grupo) > 1:
            for r in grupo:
                r['estado'] = 'incompleto'
                r['errores'].append(ERROR_EDP_DUPLICADO)
    presentes = {r['unidad'] for r in registros}
    for unidad in UNIDADES:
        if unidad not in presentes:
            avisos.append(f'{unidad}: sin archivos EDP disponibles.')
    if not registros:
        avisos.append('Agrega los Excel a Fuentes / Unidad / Año-Mes EDP N.º y pulsa Actualizar.')
    return {'actualizado': datetime.now().astimezone().isoformat(timespec='seconds'),
            'fuentes': str(raiz.resolve()), 'registros': registros, 'avisos': avisos,
            'omitidos': omitidos, 'unidades': list(UNIDADES),
            'regla': 'Monto esperado = precio por tonelada × toneladas de tickets. Andina: ítem 1.1, suma Cantidad (t) y Precio de Avance físico. El Salvador: ítem 2.5.a, suma Peso Total (kg) / 1.000 y Precio Modificación N.º 1.'}
