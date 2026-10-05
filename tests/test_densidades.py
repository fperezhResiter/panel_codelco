"""Regresiones de densidad media, ventanas, maestra y exportación desde BD."""
import json
import sqlite3
import tempfile
import threading
import unittest
from io import BytesIO
from contextlib import closing
from pathlib import Path
from unittest.mock import patch
from urllib.request import urlopen
from urllib.error import HTTPError

from openpyxl import Workbook, load_workbook
from app.base_datos import actualizar_base_datos
from app.densidades import importar_maestra, evaluar_ticket, leer_densidades, clave_material
from app.densidades_pdf import crear_pdf_densidades
from app.servidor import ServidorPanel, crear_handler
from app.residuos import estandarizar_residuo
from app.estimaciones import estimar_residuos_faltantes
try:
    from pypdf import PdfReader
except ImportError:
    PdfReader = None


def maestra(raiz, filas=None):
    libro = Workbook()
    hoja = libro.active
    hoja.title = 'Aux'
    hoja.append(['Tipo RINP', 'Dens. Inf', 'Dens. Sup(kg/m3)', 'Salvador'])
    for fila in filas or [('Madera', 100, 300), ('Cartones', 70, 150)]:
        hoja.append(fila)
    libro.save(raiz / 'Reporte de RINP por densidad.xlsx')
    libro.close()


def edp(raiz):
    carpeta = raiz / 'Codelco El Salvador' / '2026 -05 EDP 47'
    carpeta.mkdir(parents=True)
    libro = Workbook()
    hoja = libro.active
    hoja.title = 'RES.NO PELIGROSOS'
    hoja.append(['Fecha', 'N.º Ticket', 'Peso Total (kg)', 'Tipo de Residuo No Peligroso'])
    for fila in [('01/05/2026', 0, 3000, 'Madera'), ('02/05/2026', 0, 4000, 'Madera (1/2)'),
                 ('03/05/2026', 3, 100, 'Madera'), ('04/05/2026', 4, 100, 'Desconocido')]:
        hoja.append(fila)
    hoja.row_dimensions[3].hidden = True
    hoja.append(['TOTAL', None, 7200])
    hoja.append(['Resumen', 999, 999999, 'Madera'])
    libro.save(carpeta / 'EDP.xlsx')
    libro.close()


def edp_andina(raiz):
    carpeta = raiz / 'Codelco Andina' / '2026 -05 EDP 47'
    carpeta.mkdir(parents=True)
    libro = Workbook()
    hoja = libro.active
    hoja.title = '1.1'
    hoja.append(['Fecha', 'Nro. Ticket', 'Cantidad (Ton)', 'Tipo de Producto'])
    for fila in [('01/05/2026', 10, 6.175, 'RINP Granel'),
                 ('02/05/2026', 11, 9.5, 'RINP Granel (1/2)'),
                 ('03/05/2026', 12, .475, 'RINP Granel'),
                 ('04/05/2026', 13, 15, 'RINP Granel'),
                 ('05/05/2026', 14, .5, 'Polvo Roca'),
                 ('06/05/2026', 15, .5, None)]:
        hoja.append(fila)
    hoja.append(['TOTAL', None, 32.15])
    libro.save(carpeta / 'EDP.xlsx')
    libro.close()


class DensidadesTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.raiz = Path(self.temp.name)
        maestra(self.raiz)
        self.material = importar_maestra(self.raiz)['materiales'][0]

    def tearDown(self):
        self.temp.cleanup()

    def evaluar(self, volumen):
        return evaluar_ticket({'peso_kg': volumen * 200}, self.material)

    def test_promedio_y_conversion_kg_a_m3(self):
        self.assertEqual(self.material['densidad_kg_m3'], 200)
        resultado = self.evaluar(15)
        self.assertEqual(resultado['volumen_m3'], 15)
        self.assertEqual(resultado['volumen_tipo'], 15)

    def test_limites_inclusivos_y_fuera_sin_redondear(self):
        for volumen, capacidad in ((12, 15), (18, 20), (17, 15), (23, 20), (27, 30), (33, 30)):
            with self.subTest(volumen=volumen):
                resultado = self.evaluar(volumen)
                self.assertEqual(resultado['estado'], 'coincide')
                self.assertEqual(resultado['volumen_tipo'], capacidad)
        for volumen in (11.99999, 23.00001, 26.99999, 33.00001, 24, 0):
            with self.subTest(volumen=volumen):
                esperado = 'no_coincide_menor' if volumen < 15 else 'no_coincide_mayor'
                self.assertEqual(self.evaluar(volumen)['estado'], esperado)

    def test_colores_priorizan_coincidencias_y_separan_no_coincidencias(self):
        for volumen, estado in ((0, 'no_coincide_menor'), (11, 'no_coincide_menor'),
                                (12, 'coincide'), (14, 'coincide'), (15, 'coincide'),
                                (18, 'coincide'), (23, 'coincide'), (24, 'no_coincide_mayor'),
                                (27, 'coincide'), (33, 'coincide'), (34, 'no_coincide_mayor')):
            with self.subTest(volumen=volumen):
                resultado = self.evaluar(volumen)
                self.assertEqual(resultado['estado'], estado)
                if estado != 'coincide':
                    self.assertIsNone(resultado['volumen_tipo'])

    def test_andina_bandas_y_colores_independientes_de_salvador(self):
        for volumen, estado, capacidad in ((9.9999, 'no_coincide_menor', None),
                                           (10, 'coincide', 13), (13, 'coincide', 13),
                                           (16, 'coincide', 13), (16.5, 'no_coincide_mayor', None),
                                           (17, 'coincide', 20), (20, 'coincide', 20),
                                           (23, 'coincide', 20), (23.0001, 'no_coincide_mayor', None),
                                           (30, 'no_coincide_mayor', None),
                                           (36.9999, 'no_coincide_mayor', None),
                                           (37, 'coincide', 40), (40, 'coincide', 40),
                                           (43, 'coincide', 40), (43.0001, 'no_coincide_mayor', None)):
            with self.subTest(volumen=volumen):
                r = evaluar_ticket({'peso_kg': volumen * 200}, self.material, 'Andina')
                self.assertEqual(r['estado'], estado)
                self.assertEqual(r['volumen_tipo'], capacidad)
        self.assertEqual(self.evaluar(30)['volumen_tipo'], 30)
        self.assertEqual(self.evaluar(40)['estado'], 'no_coincide_mayor')

    def crear_bd_ambas_unidades(self):
        edp(self.raiz)
        edp_andina(self.raiz)
        maestra(self.raiz, [('Madera', 100, 300), ('Asimilable a RINP', 350, 600)])
        ruta = self.raiz / 'bd.sqlite3'
        actualizar_base_datos(self.raiz, ruta)
        return ruta

    def test_andina_tipo_producto_conversion_t_y_alias_solo_de_andina(self):
        ruta = self.crear_bd_ambas_unidades()
        datos = leer_densidades(ruta, {'unidad': ['Andina']})
        self.assertEqual(datos['volumenes'], [13, 20, 40])
        self.assertEqual(datos['umbral_m3'], 13)
        self.assertEqual(len(datos['tickets']), 6)
        self.assertTrue(all(t['unidad'] == 'Andina' for t in datos['tickets']))
        self.assertEqual(datos['tickets'][0]['peso_ton'], 6.175)
        self.assertEqual(datos['tickets'][0]['peso_kg'], 6175)
        self.assertEqual(datos['tickets'][0]['densidad_kg_m3'], 475)
        self.assertEqual(datos['tickets'][0]['volumen_m3'], 13)
        self.assertEqual(datos['tickets'][1]['volumen_m3'], 20)
        self.assertEqual(datos['tickets'][1]['residuo_original'], 'RINP Granel (1/2)')
        self.assertEqual([t['estado'] for t in datos['tickets']],
                         ['coincide', 'coincide', 'no_coincide_menor', 'no_coincide_mayor', 'sin_evaluar', 'no_coincide_menor'])
        self.assertTrue(any('Polvo de roca' in a for a in datos['avisos']))
        self.assertTrue(any('residuo estimado' in a for a in datos['avisos']))
        self.assertEqual(len(leer_densidades(ruta)['tickets']), 4)
        self.assertEqual(len(leer_densidades(ruta, {'unidad': ['Andina'], 'residuo': ['RINP Granel']})['tickets']), 5)
        self.assertNotEqual(estandarizar_residuo('RINP Granel', 'El Salvador'), 'Asimilable a RINP')
        with self.assertRaisesRegex(ValueError, 'unidad'):
            leer_densidades(ruta, {'unidad': ['Otra']})

    @unittest.skipUnless(PdfReader, 'Instala pypdf para verificar el contenido del reporte.')
    def test_pdf_andina_identifica_unidad_y_sus_capacidades(self):
        datos = leer_densidades(self.crear_bd_ambas_unidades(), {'unidad': ['Andina']})
        reader = PdfReader(BytesIO(crear_pdf_densidades(datos)))
        texto = '\n'.join(p.extract_text() for p in reader.pages)
        self.assertIn('Chequeo de densidades - Andina', texto)
        self.assertIn('13 m³ (±3 m³)', texto)
        self.assertIn('20 m³ (±3 m³)', texto)
        self.assertIn('40 m³ (±3 m³)', texto)
        self.assertNotIn('30 m³ (±3 m³)', texto)
        self.assertIn('RINP Granel', texto)
        self.assertIn('475,00', texto)
        self.assertIn('Polvo de roca', texto)
        self.assertIn('Tipo estimado', texto)
        self.assertNotIn('El Salvador', texto)

    def test_andina_capacidad_40_filtra_ticket_desde_bd(self):
        ruta = self.crear_bd_ambas_unidades()
        excel = self.raiz / 'Codelco Andina' / '2026 -05 EDP 47' / 'EDP.xlsx'
        libro = load_workbook(excel)
        libro['1.1']['C5'] = 19
        libro['1.1']['C8'] = 36.15
        libro.save(excel)
        libro.close()
        actualizar_base_datos(self.raiz, ruta)
        datos = leer_densidades(ruta, {'unidad': 'Andina', 'volumen_tipo': '40'})
        self.assertEqual(len(datos['tickets']), 1)
        self.assertEqual(datos['tickets'][0]['ticket'], '13')
        self.assertEqual(datos['tickets'][0]['volumen_m3'], 40)
        self.assertEqual(datos['tickets'][0]['estado'], 'coincide')

    def test_polvo_roca_usa_nueva_densidad_aux_y_no_es_estimado(self):
        ruta = self.crear_bd_ambas_unidades()
        maestra(self.raiz, [('Asimilable a RINP', 350, 600), ('Polvo de roca', 1500, 2000)])
        actualizar_base_datos(self.raiz, ruta)
        tickets = leer_densidades(ruta, {'unidad': 'Andina', 'residuo': 'Polvo Roca'})['tickets']
        self.assertEqual(len(tickets), 1)
        polvo = tickets[0]
        self.assertEqual(polvo['residuo_original'], 'Polvo Roca')
        self.assertEqual(polvo['densidad_kg_m3'], 1750)
        self.assertAlmostEqual(polvo['volumen_m3'], 500 / 1750)
        self.assertFalse(polvo['residuo_es_estimado'])

    def test_estimacion_persistida_conserva_original_y_correccion_excel_la_reemplaza(self):
        ruta = self.crear_bd_ambas_unidades()
        estimado = leer_densidades(ruta, {'unidad': 'Andina'})['tickets'][-1]
        self.assertEqual(estimado['residuo_original'], '')
        self.assertEqual(estimado['residuo'], 'Asimilable a RINP')
        self.assertEqual(estimado['densidad_kg_m3'], 475)
        self.assertTrue(estimado['residuo_es_estimado'])
        self.assertIn('4/5 tickets', estimado['criterio_estimacion_residuo'])
        actualizar_base_datos(self.raiz, ruta)
        repetido = leer_densidades(ruta, {'unidad': 'Andina'})['tickets'][-1]
        self.assertEqual(repetido['criterio_estimacion_residuo'], estimado['criterio_estimacion_residuo'])
        excel = self.raiz / 'Codelco Andina' / '2026 -05 EDP 47' / 'EDP.xlsx'
        libro = load_workbook(excel)
        libro['1.1']['D7'] = 'Madera'
        libro.save(excel)
        libro.close()
        actualizar_base_datos(self.raiz, ruta)
        corregido = leer_densidades(ruta, {'unidad': 'Andina'})['tickets'][-1]
        self.assertFalse(corregido['residuo_es_estimado'])
        self.assertEqual(corregido['residuo'], 'Madera')
        self.assertEqual(corregido['densidad_kg_m3'], 200)

    def test_estimacion_edp_cercanos_no_cruza_unidades_ni_altera_tipos_declarados(self):
        def registro(unidad, mes, tipos):
            return {'unidad': unidad, 'anio': 2026, 'mes': mes, 'periodo': f'2026-{mes:02}',
                    'edp': mes, 'archivo': f'{unidad}/{mes}.xlsx',
                    'tickets': [{'residuo': tipo, 'peso_kg': 1200} for tipo in tipos]}
        maestra(self.raiz, [('Madera', 100, 300), ('Asimilable a RINP', 350, 600)])
        registros = [registro('Andina', 1, ['RINP Granel'] * 3),
                     registro('Andina', 2, ['', 'Desconocido']),
                     registro('Andina', 3, ['RINP Granel'] * 2),
                     registro('El Salvador', 2, ['Madera'] * 9)]
        estimar_residuos_faltantes(registros, importar_maestra(self.raiz)['materiales'])
        estimado, declarado = registros[1]['tickets']
        self.assertEqual(estimado['residuo_estimado'], 'Asimilable a RINP')
        self.assertEqual(estimado['peso_kg'], 1200)
        self.assertEqual(len(estimado['referencias_estimacion_residuo']), 2)
        self.assertFalse(declarado['residuo_es_estimado'])
        self.assertEqual(declarado['residuo_original'], 'Desconocido')
        sin_evidencia = [registro('Andina', 2, ['']), registro('El Salvador', 2, ['Madera'] * 9)]
        estimar_residuos_faltantes(sin_evidencia, importar_maestra(self.raiz)['materiales'])
        self.assertFalse(sin_evidencia[0]['tickets'][0]['residuo_es_estimado'])

    def test_solapamiento_y_empate_mantienen_todas_las_compatibilidades(self):
        resultado = self.evaluar(17.5)
        self.assertEqual(resultado['volumen_tipo'], 15)
        self.assertEqual(resultado['capacidades_compatibles'], [15, 20])
        self.assertEqual(self.evaluar(17.51)['volumen_tipo'], 20)

    def test_invalidos_no_se_convierten_en_rojos_ni_ceros(self):
        for peso in (None, -10, 'NaN', True):
            self.assertEqual(evaluar_ticket({'peso_kg': peso}, self.material)['estado'], 'sin_evaluar')
        self.assertEqual(evaluar_ticket({'peso_kg': 3000}, None)['estado'], 'sin_evaluar')

    def test_normalizacion_conservadora_y_fracciones(self):
        self.assertEqual(clave_material('  PLÁSTICO (1/2)  '), clave_material('Plastico'))
        self.assertEqual(clave_material('Aceite'), clave_material('Aceite en desuso'))
        self.assertNotEqual(clave_material('Madera contaminada'), clave_material('Madera'))

    def test_variantes_se_agrupan_sin_perder_el_nombre_y_peso_original(self):
        grupos = {
            'Asimilable a RINP': ['Asimable a RINP', 'ASimable a RINP', 'Asimable a RINP (1/2)',
                                 'Asimilable a RINP (2/2)', 'ASIMILABLES A RINP'],
            'Madera': ['MADERA', 'Madera (1/3)', 'Maderas'],
            'Chatarra': ['Chatarra (1/2)', 'Chatarra /2/2)', 'Chatarras'],
            'Aceite en desuso': ['Aceite  ', 'Aceites en desuso'],
            'Gomas en desuso': ['Goma', 'Gomas (1 / 2)'],
            'Plástico': ['PLASTICO', 'Plásticos (2/3)'],
            'Cartones': ['Cartón', 'CARTONES (1/2)'],
            'Filtros en desuso': ['Filtro en desuso', 'Filtros (1/2)'],
        }
        for nombre, variantes in grupos.items():
            for variante in variantes:
                with self.subTest(variante=variante):
                    self.assertEqual(estandarizar_residuo(variante), nombre)
                    self.assertEqual(clave_material(variante), clave_material(nombre))
                    r = evaluar_ticket({'residuo': variante, 'peso_kg': 3000}, self.material)
                    self.assertEqual(r['residuo'], nombre)
                    self.assertEqual(r['residuo_original'], variante)
                    self.assertEqual(r['peso_kg'], 3000)
                    self.assertEqual(r['volumen_m3'], 15)

    def test_bd_y_filtro_agrupan_variantes_y_conservan_originales(self):
        maestra(self.raiz, [('Asimilable a RINP', 100, 300)])
        edp(self.raiz)
        ruta_excel = self.raiz / 'Codelco El Salvador' / '2026 -05 EDP 47' / 'EDP.xlsx'
        libro = load_workbook(ruta_excel)
        variantes = ['Asimable a RINP', 'ASimable a RINP', 'Asimable a RINP (1/2)']
        for fila, variante in enumerate(variantes, 2):
            libro.active.cell(fila, 4, variante)
        libro.save(ruta_excel)
        libro.close()
        ruta_bd = self.raiz / 'bd.sqlite3'
        actualizar_base_datos(self.raiz, ruta_bd)
        datos = leer_densidades(ruta_bd, {'residuo': ['Asimilable a RINP']})
        self.assertEqual(len(datos['tickets']), 3)
        self.assertEqual({t['residuo'] for t in datos['tickets']}, {'Asimilable a RINP'})
        self.assertEqual([t['residuo_original'] for t in datos['tickets']], variantes)
        self.assertTrue(all(t['densidad_kg_m3'] == 200 for t in datos['tickets']))
        self.assertEqual(len(leer_densidades(ruta_bd, {'residuo': ['Asimable a RINP (1/2)']})['tickets']), 3)
        with closing(sqlite3.connect(ruta_bd)) as conexion:
            ticket = json.loads(conexion.execute('SELECT datos FROM tickets ORDER BY orden LIMIT 1').fetchone()[0])
        self.assertEqual(ticket['residuo_estandarizado'], 'Asimilable a RINP')
        self.assertEqual(ticket['residuo_original'], variantes[0])

    def test_solo_aux_aunque_otras_hojas_tengan_tablas_de_densidad(self):
        ruta = self.raiz / 'Reporte de RINP por densidad.xlsx'
        libro = load_workbook(ruta)
        otra = libro.create_sheet('Otra tabla')
        otra.append(['Tipo RINP', 'Dens. Inf', 'Dens. Sup(kg/m3)'])
        otra.append(['Madera', 900, 1000])
        libro.save(ruta)
        datos = importar_maestra(self.raiz)
        self.assertEqual(datos['materiales'][0]['densidad_kg_m3'], 200)
        self.assertTrue(all(m['hoja'] == 'Aux' for m in datos['materiales']))
        libro.remove(libro['Aux'])
        libro.save(ruta)
        libro.close()
        datos = importar_maestra(self.raiz)
        self.assertEqual(datos['materiales'], [])
        self.assertTrue(any('hoja AUX' in a for a in datos['avisos']))

    def test_maestra_duplicada_y_densidad_invalida_no_eligen_un_valor(self):
        maestra(self.raiz, [('Madera', 100, 300), ('MADERA', 100, 400), ('Cartones', 150, 70), ('Goma', 0, 10)])
        datos = importar_maestra(self.raiz)
        self.assertEqual(len(datos['materiales']), 3)
        self.assertTrue(all(m['error'] and m['densidad_kg_m3'] is None for m in datos['materiales']))

    def test_formula_sin_cache_y_falta_maestra(self):
        maestra(self.raiz, [('Madera', '=100+10', 300)])
        datos = importar_maestra(self.raiz)
        self.assertIn('fórmula sin resultado', datos['materiales'][0]['error'])
        (self.raiz / 'Reporte de RINP por densidad.xlsx').unlink()
        self.assertTrue(importar_maestra(self.raiz)['avisos'])

    def crear_bd(self):
        edp(self.raiz)
        ruta = self.raiz / 'bd.sqlite3'
        actualizar_base_datos(self.raiz, ruta)
        return ruta

    def test_bd_guarda_tickets_aunque_falte_imputacion_y_no_relee_excel(self):
        ruta = self.crear_bd()
        with patch('app.densidades.FuenteExcel', side_effect=AssertionError('No debe leer Excel')):
            datos = leer_densidades(ruta)
        self.assertEqual(len(datos['tickets']), 4)
        self.assertEqual([t['ticket'] for t in datos['tickets']], ['0', '0', '3', '4'])
        self.assertEqual([t['estado'] for t in datos['tickets']], ['coincide', 'coincide', 'no_coincide_menor', 'sin_evaluar'])
        self.assertEqual(len(datos['maestra']['materiales']), 2)

    def test_cambio_y_eliminacion_maestra_se_reflejan_sin_cambiar_tickets(self):
        ruta = self.crear_bd()
        maestra(self.raiz, [('Madera', 400, 600)])
        actualizar_base_datos(self.raiz, ruta)
        self.assertEqual(leer_densidades(ruta)['tickets'][0]['volumen_m3'], 6)
        (self.raiz / 'Reporte de RINP por densidad.xlsx').unlink()
        actualizar_base_datos(self.raiz, ruta)
        self.assertTrue(all(t['estado'] == 'sin_evaluar' for t in leer_densidades(ruta)['tickets']))

    def test_filtros_pdf_y_busqueda_ticket_cero(self):
        ruta = self.crear_bd()
        datos = leer_densidades(ruta, {'anio': ['2026'], 'mes': ['5'], 'edp': ['47'],
                                      'estado': ['coincide'], 'volumen_tipo': ['15'], 'buscar': ['0']})
        self.assertEqual(len(datos['tickets']), 1)
        self.assertTrue(crear_pdf_densidades(datos).startswith(b'%PDF-'))
        with self.assertRaisesRegex(ValueError, 'No hay tickets'):
            crear_pdf_densidades(leer_densidades(ruta, {'edp': ['999']}))

    def test_bd_anterior_pide_actualizacion(self):
        ruta = self.crear_bd()
        with closing(sqlite3.connect(ruta)) as conexion:
            with conexion:
                conexion.execute('DROP TABLE maestra_densidades')
        with self.assertRaisesRegex(ValueError, 'Actualizar_BD'):
            leer_densidades(ruta)

    def test_http_json_pdf_filtrado_y_js(self):
        ruta = self.crear_bd_ambas_unidades()
        servidor = ServidorPanel(('127.0.0.1', 0), crear_handler(self.raiz, self.raiz / 'auditoria', ruta))
        hilo = threading.Thread(target=servidor.serve_forever, daemon=True)
        hilo.start()
        base = f'http://127.0.0.1:{servidor.server_port}'
        try:
            with patch('app.densidades.FuenteExcel', side_effect=AssertionError('No debe leer Excel')):
                with urlopen(base + '/api/densidades?estado=no_coincide_menor') as respuesta:
                    self.assertEqual(len(json.load(respuesta)['tickets']), 1)
                with urlopen(base + '/api/densidades/reporte.pdf?estado=coincide') as respuesta:
                    self.assertEqual(respuesta.headers['Content-Type'], 'application/pdf')
                    self.assertIn('attachment', respuesta.headers['Content-Disposition'])
                    self.assertTrue(respuesta.read().startswith(b'%PDF-'))
                with urlopen(base + '/web/js/densidades.js') as respuesta:
                    self.assertIn(b'/api/densidades', respuesta.read())
                with urlopen(base + '/api/densidades?unidad=Andina&estado=sin_evaluar') as respuesta:
                    datos = json.load(respuesta)
                    self.assertEqual(datos['volumenes'], [13, 20, 40])
                    self.assertEqual(len(datos['tickets']), 1)
                with urlopen(base + '/api/densidades/reporte.pdf?unidad=Andina&estado=coincide') as respuesta:
                    self.assertIn('andina', respuesta.headers['Content-Disposition'])
                    self.assertTrue(respuesta.read().startswith(b'%PDF-'))
                with self.assertRaises(HTTPError) as error:
                    urlopen(base + '/api/densidades/reporte.pdf?edp=999')
                self.assertEqual(error.exception.code, 400)
        finally:
            servidor.shutdown()
            servidor.server_close()
            hilo.join()


if __name__ == '__main__':
    unittest.main()
