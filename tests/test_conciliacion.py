"""Regresiones de la regla de negocio y de la lectura de tablas."""
import json
import os
import tempfile
import threading
import unittest
from datetime import datetime
from decimal import Decimal
from pathlib import Path
from unittest.mock import patch
from urllib.error import HTTPError
from urllib.request import urlopen
from openpyxl import Workbook
from openpyxl.worksheet.table import Table
from app.excel import FuenteExcel, numero
from app.reportes import comparar, leer_imputacion, leer_tickets, metadatos, crear_reporte
from app.andina import leer_avance_andina
from app.base_datos import actualizar_base_datos, leer_reporte
from app.servidor import ServidorPanel, crear_handler


def fuente():
    libro = FuenteExcel.__new__(FuenteExcel)
    libro.valores = Workbook()
    libro.valores.remove(libro.valores.active)
    detalle = libro.valores.create_sheet('DETALLE DE IMPUTACIONES N° 47')
    detalle.merge_cells('G1:I1'); detalle['G1'] = 'MODIFICACIÓN N° 1'
    detalle.merge_cells('J1:L1'); detalle['J1'] = 'AVANCE FISICO DE LAS PARTIDAS'
    detalle.merge_cells('M1:O1'); detalle['M1'] = 'AVANCE FINANCIERO'
    detalle.merge_cells('H2:I2'); detalle['H2'] = 'Precio'
    for celda, valor in {'G3':'Cantidad','H3':'Precio','I3':'Total','K3':'Anterior EP N° 46',
                         'L3':'EP N°47','N3':'Anterior EP N° 46','O3':'EP N°47',
                         'A5':'2.5.a','E5':999,'H5':100,'I5':800,'L5':3,'O5':300}.items():
        detalle[celda] = valor
    hoja = libro.valores.create_sheet('RES.NO PELIGROSOS')
    for c,v in {'B2':'Fecha','C2':'N° TICKET','D2':'Peso Total (Kg)',
                'B3':datetime(2026,4,21),'C3':0,'D3':1000,
                'B4':'13-05-2026','C4':0,'D4':2000,
                'C5':'TOTAL KG','D5':3000,'C6':'TOTAL TON','D6':3,
                'C8':'Resumen','D8':99999}.items():
        hoja[c]=v
    hoja.add_table(Table(displayName='Tickets', ref='B2:D4'))
    libro.formulas = Workbook()
    libro.formulas.remove(libro.formulas.active)
    for s in libro.valores:
        f=libro.formulas.create_sheet(s.title)
        for row in s:
            for c in row:
                if c.value is not None:
                    f[c.coordinate]=c.value
    return libro


class ConciliacionTests(unittest.TestCase):
    def setUp(self):
        self.libro = fuente()

    def tearDown(self):
        self.libro.cerrar()

    def test_encabezados_no_toman_precio_base_ni_acumulados(self):
        valores, refs, _ = leer_imputacion(self.libro,47)
        self.assertEqual(valores, {'precio':Decimal(100),'ton_edp':Decimal(3),'monto_edp':Decimal(300)})
        self.assertEqual(refs['precio']['celda'],'H5')
        self.assertEqual(refs['monto_edp']['celda'],'O5')
        with self.assertRaisesRegex(ValueError,'47|48'):
            leer_imputacion(self.libro,48)

    def test_suma_solo_tabla_incluye_cero_repetidos_y_fecha_texto(self):
        peso,tickets,ref,errores,avisos = leer_tickets(self.libro)
        self.assertEqual(peso,Decimal(3000))
        self.assertEqual(len(tickets),2)
        self.assertEqual(ref['rango'],'D3:D4')
        self.assertEqual(tickets[1]['fecha'],'2026-05-13')
        self.assertFalse(errores)
        self.assertTrue(any('repetidos' in a for a in avisos))

    def test_sin_tabla_termina_antes_del_total_y_resumen(self):
        del self.libro.valores['RES.NO PELIGROSOS'].tables['Tickets']
        self.assertEqual(leer_tickets(self.libro)[0],Decimal(3000))

    def test_peso_faltante_no_pasa_como_cero(self):
        self.libro.valores['RES.NO PELIGROSOS']['D4']=None
        peso,_,_,errores,_ = leer_tickets(self.libro)
        self.assertIsNone(peso)
        self.assertTrue(errores)

    def test_cero_es_un_peso_valido(self):
        self.libro.valores['RES.NO PELIGROSOS']['D4']=0
        self.assertEqual(leer_tickets(self.libro)[0],Decimal(1000))

    def test_formula_sin_cache_pide_recalcular(self):
        hoja=self.libro.valores['RES.NO PELIGROSOS']
        hoja['D4']=None
        self.libro.formulas[hoja.title]['D4']='=G4-H4'
        self.assertIn('fórmula sin resultado',leer_tickets(self.libro)[3][0])

    def test_falta_ticket_bloquea_total(self):
        self.libro.valores['RES.NO PELIGROSOS']['C4']=None
        self.assertIsNone(leer_tickets(self.libro)[0])

    def test_monetario_y_toneladas_se_comparan_independientemente(self):
        correcto = comparar(Decimal(100),Decimal(3),Decimal('300.000000006'),Decimal(3000))
        self.assertEqual(correcto['estado'],'cuadra')
        self.assertEqual(comparar(Decimal(100),Decimal(3),Decimal(301),Decimal(3000))['diferencia'],1)
        self.assertEqual(comparar(Decimal(100),Decimal(4),Decimal(300),Decimal(3000))['estado'],'diferencia')

    def test_ambiguedad_de_precio_se_rechaza(self):
        self.libro.valores.worksheets[0]['G3']='Precio'
        with self.assertRaisesRegex(ValueError,'ambigua'):
            leer_imputacion(self.libro,47)

    def test_ruta_y_numeros_chilenos(self):
        r=metadatos(Path('/Fuentes/Codelco Andina/2026 -05 EDP 47/libro.xlsx'),Path('/Fuentes'))
        self.assertEqual((r['unidad'],r['periodo'],r['edp']),('Andina','2026-05',47))
        self.assertEqual(numero('141.045,50'),Decimal('141045.50'))
        self.assertEqual(numero(0),Decimal(0))
        for v in (None,'texto','NaN',True):
            with self.assertRaises(ValueError):
                numero(v)

    def test_fuentes_ausentes_no_inventa_registros(self):
        with tempfile.TemporaryDirectory() as temp:
            r=crear_reporte(Path(temp)/'ausente')
        self.assertFalse(r['registros'])
        self.assertTrue(r['avisos'])

    def test_dos_versiones_mismo_edp_no_se_consolidan(self):
        with tempfile.TemporaryDirectory() as temp:
            carpeta=Path(temp)/'Codelco El Salvador'/'2026 -05 EDP 47'
            carpeta.mkdir(parents=True)
            (carpeta/'v1.xlsx').touch(); (carpeta/'v2.xlsx').touch()
            with patch('app.reportes.FuenteExcel',side_effect=lambda _:fuente()):
                r=crear_reporte(Path(temp))
            self.assertEqual(len(r['registros']),2)
            self.assertTrue(all(x['estado']=='incompleto' for x in r['registros']))

    def test_http_publica_panel_y_no_expone_fuentes(self):
        with tempfile.TemporaryDirectory() as temp:
            base_datos = Path(temp) / 'conciliacion.sqlite3'
            actualizar_base_datos(Path(temp) / 'Fuentes', base_datos)
            servidor=ServidorPanel(('127.0.0.1',0),crear_handler(Path(temp), base_datos=base_datos))
            hilo=threading.Thread(target=servidor.serve_forever,daemon=True); hilo.start()
            base=f'http://127.0.0.1:{servidor.server_port}'
            try:
                with urlopen(base+'/') as respuesta:
                    self.assertIn(b'Estados de pago',respuesta.read())
                with urlopen(base+'/api/conciliacion') as respuesta:
                    self.assertEqual(json.load(respuesta)['registros'],[])
                with self.assertRaises(HTTPError) as error:
                    urlopen(base+'/Fuentes/secreto.xlsx')
                self.assertEqual(error.exception.code,404)
            finally:
                servidor.shutdown(); servidor.server_close(); hilo.join()


def fuente_andina(tabla=True):
    libro = FuenteExcel.__new__(FuenteExcel)
    libro.valores = Workbook()
    fisico = libro.valores.active
    fisico.title = 'Avance Físico'
    fisico.append(['Posición', 'Total Valor Presupuesto', 'Total Valor Actual EP', 'Precio', 'Monto EDP'])
    fisico.append(['1.1', 9999, 3, 100, 99999])
    financiero = libro.valores.create_sheet('Avance Financiero')
    financiero.append(['Posición', 'Valor Total Presupuesto', 'Total Valor Actual EP', 'Saldo'])
    financiero.append(['1.1', 99999, 300, 8000])
    tickets = libro.valores.create_sheet('1.1')
    tickets.append(['Fecha', 'Nro. Ticket', 'Cantidad (Ton)', 'Precio x ton', 'Valor Pesaje', 'Nr. EDP', 'Tipo Residuo'])
    tickets.append([datetime(2026, 4, 26), 0, 1, 100, 100, 55, 'RINP'])
    tickets.append([datetime(2026, 5, 25), 0, 2, 100, 200, 55, 'RINP'])
    tickets.append([None, None, None, 100, 0, 55, 'RINP'])
    tickets.row_dimensions[3].hidden = True
    if tabla:
        tickets.add_table(Table(displayName='TicketsAndina', ref='A1:G4'))
    tickets.append(['TOTAL', None, 3, None, 300])
    tickets.append(['Resumen', 999, 99999])
    libro.formulas = Workbook()
    libro.formulas.remove(libro.formulas.active)
    for hoja in libro.valores:
        formulas = libro.formulas.create_sheet(hoja.title)
        for fila in hoja:
            for c in fila:
                if c.value is not None:
                    formulas[c.coordinate] = c.value
    return libro


class AndinaTests(unittest.TestCase):
    def setUp(self):
        self.libro = fuente_andina()

    def tearDown(self):
        self.libro.cerrar()

    def test_precio_y_totales_se_leen_de_las_hojas_correctas(self):
        valores, refs, _ = leer_avance_andina(self.libro)
        self.assertEqual(valores, {'precio': Decimal(100), 'ton_edp': Decimal(3), 'monto_edp': Decimal(300)})
        self.assertEqual(refs['precio']['hoja'], 'Avance Físico')
        self.assertEqual(refs['ton_edp']['celda'], 'C2')
        self.assertEqual(refs['monto_edp']['hoja'], 'Avance Financiero')

    def test_cantidad_ya_es_toneladas_y_no_cuenta_plantillas(self):
        peso, tickets, ref, errores, avisos = leer_tickets(self.libro, 'Andina')
        self.assertEqual(peso, Decimal(3000))
        self.assertEqual([t['peso_ton'] for t in tickets], [1, 2])
        self.assertEqual(len(tickets), 2)
        self.assertEqual(ref['unidad'], 't')
        self.assertEqual(ref['hoja'], '1.1')
        self.assertFalse(errores)
        self.assertTrue(any('repetidos' in a for a in avisos))
        self.assertEqual(comparar(Decimal(100), Decimal(3), Decimal(300), peso)['estado'], 'cuadra')

    def test_formato_antiguo_sin_tabla_y_total_sin_etiqueta(self):
        hoja = self.libro.valores['1.1']
        del hoja.tables['TicketsAndina']
        hoja['C1'] = 'CANTIDAD'
        hoja['A5'] = None
        self.libro.formulas['1.1']['C5'] = '=SUM(C2:C4)'
        self.assertEqual(leer_tickets(self.libro, 'Andina')[0], Decimal(3000))

    def test_total_dentro_de_tabla_no_se_suma_dos_veces(self):
        self.libro.valores['1.1'].tables['TicketsAndina'].ref = 'A1:G5'
        self.assertEqual(leer_tickets(self.libro, 'Andina')[0], Decimal(3000))

    def test_ticket_con_cantidad_faltante_no_es_plantilla(self):
        self.libro.valores['1.1']['C3'] = None
        peso, tickets, _, errores, _ = leer_tickets(self.libro, 'Andina')
        self.assertIsNone(peso)
        self.assertEqual(len(tickets), 2)
        self.assertTrue(errores)

    def test_ticket_faltante_con_suma_de_otra_columna_no_es_total(self):
        self.libro.valores['1.1']['B3'] = None
        self.libro.formulas['1.1']['C3'] = '=SUM(D3:F3)'
        peso, tickets, _, errores, _ = leer_tickets(self.libro, 'Andina')
        self.assertIsNone(peso)
        self.assertEqual(len(tickets), 2)
        self.assertTrue(any('falta el número de ticket' in e for e in errores))

    def test_errores_y_ambiguedades_no_inventan_resultados(self):
        self.libro.valores['Avance Físico']['D2'] = None
        self.libro.formulas['Avance Físico']['D2'] = '=OtraHoja!A1'
        with self.assertRaisesRegex(ValueError, 'fórmula sin resultado'):
            leer_avance_andina(self.libro)
        self.libro.valores['Avance Físico']['F1'] = 'Precio'
        with self.assertRaisesRegex(ValueError, 'columna única'):
            leer_avance_andina(self.libro)

    def test_actualiza_andina_con_regla_nueva_aunque_el_excel_no_cambie(self):
        with tempfile.TemporaryDirectory() as temp:
            raiz = Path(temp)
            carpeta = raiz / 'Fuentes' / 'Codelco Andina' / '2026 -05 EDP 55'
            carpeta.mkdir(parents=True)
            ruta = carpeta / 'EDP.xlsx'
            ruta.touch()
            stat = ruta.stat()
            relativa = ruta.relative_to(raiz / 'Fuentes').as_posix()
            antiguo = {'unidad': 'Andina', 'estado': 'incompleto', 'errores': ['Regla antigua']}
            with patch('app.reportes.FuenteExcel', return_value=self.libro) as abrir:
                reporte = crear_reporte(raiz / 'Fuentes', {relativa: ((stat.st_size, stat.st_mtime_ns), antiguo)},
                                        {relativa: (stat.st_size, stat.st_mtime_ns)})
            abrir.assert_called_once()
            registro = reporte['registros'][0]
            self.assertEqual(registro['estado'], 'cuadra', registro['errores'])
            self.assertEqual(registro['item'], '1.1')
            self.assertEqual(registro['ton_tickets'], 3)
            with patch('app.reportes.FuenteExcel', return_value=fuente_andina()):
                actualizar_base_datos(raiz / 'Fuentes', raiz / 'reporte.sqlite3')
            self.assertEqual(leer_reporte(raiz / 'reporte.sqlite3')['registros'][0]['estado'], 'cuadra')
            with patch('app.reportes.FuenteExcel', side_effect=AssertionError('Debe reutilizar la regla nueva')):
                actualizado = actualizar_base_datos(raiz / 'Fuentes', raiz / 'reporte.sqlite3')
            self.assertEqual(actualizado['registros'][0]['ton_tickets'], 3)


@unittest.skipUnless(os.environ.get('PRUEBA_FUENTES'),'Configura PRUEBA_FUENTES para verificar los Excel reales.')
class FuentesRealesTests(unittest.TestCase):
    def test_cuatro_edp_y_suma_independiente(self):
        reporte=crear_reporte(Path(os.environ['PRUEBA_FUENTES']))
        esperado={47:(129,373200),48:(157,468220),49:(149,405950),50:(123,362070)}
        # Esta referencia histórica cubre mayo-agosto 2026 de El Salvador;
        # otras unidades y períodos agregados a Fuentes no pertenecen a su suma.
        seleccion=[r for r in reporte['registros'] if r['unidad']=='El Salvador' and
                   r['anio']==2026 and r['edp'] in esperado]
        self.assertEqual(len(seleccion),len(esperado))
        registros={r['edp']:r for r in seleccion}
        for edp,(cantidad,kg) in esperado.items():
            r=registros[edp]
            self.assertEqual(r['estado'],'cuadra',r['errores'])
            self.assertEqual(r['peso_kg'],kg)
            self.assertEqual(len(r['tickets']),cantidad)
            self.assertEqual(sum(Decimal(str(t['peso_kg'])) for t in r['tickets']),Decimal(kg))
            self.assertEqual(r['precio'],141045)
            self.assertEqual(r['referencias']['precio']['celda'],'H47')
            self.assertEqual(r['referencias']['monto_edp']['celda'],'O47')
        self.assertEqual(sum(r['ton_tickets'] for r in registros.values()),1609.44)


if __name__ == '__main__':
    unittest.main()
