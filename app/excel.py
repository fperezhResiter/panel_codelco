
import re
import unicodedata
import warnings
from datetime import date, datetime
from decimal import Decimal, InvalidOperation
from openpyxl import load_workbook


def normalizar(valor):
    texto = ' '.join(str(valor if valor is not None else '').split()).upper()
    return ''.join(c for c in unicodedata.normalize('NFD', texto) if not unicodedata.combining(c))


def clave(valor):
    return re.sub(r'[^A-Z0-9]', '', normalizar(valor))


def numero(valor):
    if valor is None or isinstance(valor, bool):
        raise ValueError('valor numérico vacío o inválido')
    if isinstance(valor, str):
        valor = valor.strip().replace('\u00a0', '').replace(' ', '').replace('$', '')
        if ',' in valor:
            valor = valor.replace('.', '').replace(',', '.')
        elif re.fullmatch(r'-?\d{1,3}(\.\d{3})+', valor):
            valor = valor.replace('.', '')
    try:
        resultado = Decimal(str(valor))
    except InvalidOperation as exc:
        raise ValueError('valor no numérico') from exc
    if not resultado.is_finite():
        raise ValueError('valor no finito')
    return resultado


def fecha(valor):
    if isinstance(valor, datetime):
        return valor.date().isoformat()
    if isinstance(valor, date):
        return valor.isoformat()
    for formato in ('%d-%m-%Y', '%d/%m/%Y', '%Y-%m-%d'):
        try:
            return datetime.strptime(str(valor).strip(), formato).date().isoformat()
        except ValueError:
            pass
    return None


class FuenteExcel:
    def __init__(self, ruta):
        self.ruta = ruta
        # Estas advertencias se refieren al guardado, que esta aplicación no realiza.
        with warnings.catch_warnings():
            warnings.filterwarnings('ignore', message='DrawingML support is incomplete.*')
            warnings.filterwarnings('ignore', message='Data Validation extension is not supported.*')
            self.valores = load_workbook(ruta, data_only=True)
            try:
                self.formulas = load_workbook(ruta, data_only=False)
            except Exception:
                self.valores.close()
                raise

    def cerrar(self):
        self.valores.close()
        self.formulas.close()

    def hoja(self, patron):
        candidatas = [h for h in self.valores if re.fullmatch(patron, clave(h.title))]
        if len(candidatas) != 1:
            raise ValueError(f'Se esperaba una hoja compatible con {patron}; encontradas: {len(candidatas)}.')
        return candidatas[0]

    def decimal(self, hoja, fila, columna):
        celda = hoja.cell(fila, columna)
        original = self.formulas[hoja.title].cell(fila, columna)
        referencia = f"'{hoja.title}'!{celda.coordinate}"
        if original.data_type == 'f' and celda.value is None:
            raise ValueError(f'{referencia}: fórmula sin resultado guardado. Recalcula y guarda el archivo en Excel.')
        try:
            return numero(celda.value)
        except ValueError as exc:
            raise ValueError(f'{referencia}: {exc}.') from exc

    def referencia(self, hoja, fila, columna):
        celda = self.formulas[hoja.title].cell(fila, columna)
        return {'hoja': hoja.title, 'celda': celda.coordinate,
                'formula': celda.value if celda.data_type == 'f' else None}


def extension(hoja, celda):
    for rango in hoja.merged_cells.ranges:
        if celda.coordinate in rango:
            return rango.min_col, rango.max_col, rango.max_row
    return celda.column, celda.column, celda.row


def unico(candidatos, mensaje):
    if len(candidatos) != 1:
        raise ValueError(mensaje)
    return candidatos[0]


def columna_precio(hoja, grupo, fila_item):
    izquierda, derecha, ultima = extension(hoja, grupo)
    candidatos = []
    for fila in hoja.iter_rows(min_row=ultima + 1, max_row=fila_item - 1,
                               min_col=izquierda, max_col=derecha):
        for celda in fila:
            if clave(celda.value) in ('PRECIO', 'PRECIOUNITARIO'):
                inicio, fin, _ = extension(hoja, celda)
                if inicio == fin:
                    candidatos.append(celda)
    if not candidatos:
        raise ValueError('Falta la columna Precio de Modificación N.º 1.')
    inferior = max(c.row for c in candidatos)
    return unico([c.column for c in candidatos if c.row == inferior],
                 'La columna Precio de Modificación N.º 1 es ambigua.')


def columna_edp(hoja, grupo, fila_item, edp):
    izquierda, derecha, ultima = extension(hoja, grupo)
    candidatos = []
    for fila in hoja.iter_rows(min_row=ultima + 1, max_row=fila_item - 1,
                               min_col=izquierda, max_col=derecha):
        for celda in fila:
            match = re.fullmatch(r'(?:EDP|EP)(?:N|NO|NUMERO)?0*(\d+)', clave(celda.value))
            if match and int(match[1]) == edp:
                candidatos.append(celda.column)
    return unico(candidatos, f'No se identifica un único EP/EDP N.º {edp} en {grupo.value}.')
