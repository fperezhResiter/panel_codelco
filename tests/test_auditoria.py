"""Pruebas del muestreo, persistencia y exportaciones de la auditoría manual."""
import base64
import csv
import json
import tempfile
import threading
import unittest
from concurrent.futures import ThreadPoolExecutor
from datetime import datetime
from io import BytesIO, StringIO
from pathlib import Path
from urllib.error import HTTPError
from urllib.request import Request, urlopen
from urllib.parse import urlsplit
from openpyxl import Workbook
from openpyxl.worksheet.table import Table
from app.auditoria import Auditoria, crear_csv, leer_excel, resumen_edp, MAX_ARCHIVO, MAX_SOLICITUD, decodificar, EXT_DOCUMENTOS
from app.auditoria_pdf import crear_pdf_auditoria, crear_pdf_unidad
from app.servidor import ServidorPanel, crear_handler


def excel(cantidad=8, tabla=True, hoja='RES.NO PELIGROSO'):
    libro = Workbook()
    h = libro.active
    h.title = hoja
    h.append(['Fecha', 'N.º Ticket', 'Conductor', 'Peso Total (kg)', 'OBS', 'Campo adicional'])
    for n in range(cantidad):
        h.append([datetime(2026, 4, 25), 0 if n < 2 else n, 'Ana & José', n * 120,
                  '=1+2' if n == 0 else 'Texto largo; con "comillas"\ny salto', 'Dato adicional'])
    h.row_dimensions[3].hidden = True
    h['B4'].number_format = '000000'
    if tabla and cantidad:
        h.add_table(Table(displayName='Tickets', ref=f'A1:F{cantidad+1}'))
    h.append(['TOTAL', None, None, 99999])
    h.append(['Resumen inferior', 999, None, 99999])
    salida = BytesIO(); libro.save(salida); libro.close()
    return salida.getvalue()


def adjunto(nombre='respaldo.pdf', contenido=b'%PDF-1.4\nRespaldo de prueba'):
    return {'nombre': nombre, 'contenido': base64.b64encode(contenido).decode()}


def excel_andina(tabla=True, total_sin_etiqueta=False, cantidad_vacia=False):
    libro = Workbook()
    hoja = libro.active
    hoja.title = '1.1'
    hoja.append(['Fecha', 'Nro. Ticket', 'Cantidad', 'Precio x ton', 'Valor Pesaje', 'Nr. EDP'])
    for n in range(8):
        hoja.append([datetime(2026, 4, 26), 0 if n < 2 else n, (n + 1) / 10, 56351, '=C2*D2', 55])
    hoja['B4'].number_format = '000000'
    hoja.row_dimensions[3].hidden = True
    if cantidad_vacia:
        hoja['C3'] = None
    for _ in range(2):
        hoja.append([None, None, None, 56351, '=C10*D10', 55])
    if tabla:
        hoja.add_table(Table(displayName='TicketsAndina', ref='A1:F11'))
    hoja.append([None if total_sin_etiqueta else 'TOTAL', None, '=SUM(C2:C11)'])
    hoja.append(['Resumen posterior', 999, 99999])
    salida = BytesIO()
    libro.save(salida)
    libro.close()
    return salida.getvalue()


class AuditoriaTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.raiz = Path(self.temp.name)
        self.fuentes = self.raiz / 'Fuentes'
        carpeta = self.fuentes / 'Codelco El Salvador' / '2026 -05 EDP 47'
        carpeta.mkdir(parents=True)
        self.archivo = carpeta / 'EDP.xlsx'
        self.archivo.write_bytes(excel())
        self.a = Auditoria(self.fuentes, self.raiz / 'Datos')
        self.fuente = self.a.catalogo()['fuentes'][0]['id']

    def tearDown(self):
        self.temp.cleanup()

    def sortear(self, cantidad=3):
        return self.a.sortear({'fuente': self.fuente, 'cantidad': cantidad})

    def seleccionar(self, fila=3):
        lista = self.a.listar_tickets(self.fuente)
        return self.a.seleccionar({'fuente': self.fuente, 'fila': fila, 'sha256': lista['sha256']})

    def agregar_andina(self):
        carpeta = self.fuentes / 'Codelco Andina' / '2026 -05 EDP 55'
        carpeta.mkdir(parents=True)
        archivo = carpeta / 'Andina.xlsx'
        archivo.write_bytes(excel_andina())
        fuente = next(f for f in self.a.catalogo()['fuentes'] if f['unidad'] == 'Andina')
        return fuente['id'], archivo

    def test_andina_lectura_toneladas_y_ambos_formatos(self):
        for tabla in (True, False):
            for sin_etiqueta in (True, False):
                datos = leer_excel(excel_andina(tabla, sin_etiqueta), 'Andina')
                self.assertEqual(datos['hoja'], '1.1')
                self.assertEqual(datos['item'], '1.1')
                self.assertEqual(len(datos['tickets']), 8)
                self.assertEqual(datos['tickets'][1]['fila'], 3)
                self.assertEqual(datos['tickets'][2]['ticket'], '000002')
                self.assertEqual(datos['tickets'][1]['campos'][2]['valor'], '0.2')
                self.assertEqual(datos['tickets'][1]['campos'][2]['unidad'], 't')
                self.assertTrue(any('repetidos' in a for a in datos['avisos']))
        datos = leer_excel(excel_andina(cantidad_vacia=True), 'Andina')
        self.assertEqual(len(datos['tickets']), 8)
        self.assertEqual(datos['tickets'][1]['campos'][2]['valor'], '')

    def test_andina_sortea_3_4_5_y_conserva_historial(self):
        fuente, _ = self.agregar_andina()
        for cantidad in (3, 4, 5):
            r = self.solicitud_memoria('POST', '/api/auditoria/sortear', {'fuente': fuente, 'cantidad': cantidad})
            self.assertEqual(r['status'], 200)
            muestra = json.loads(r['contenido'])
            self.assertEqual(muestra['unidad'], 'Andina')
            self.assertEqual(muestra['hoja'], '1.1')
            self.assertEqual(muestra['poblacion'], 8)
            self.assertEqual(len({t['fila'] for t in muestra['tickets']}), cantidad)
            self.assertTrue(all(t['fila'] < 10 for t in muestra['tickets']))
            self.assertEqual(self.a.obtener(muestra['id'])['tickets'], muestra['tickets'])
        self.assertEqual(len(self.a.catalogo()['muestras']), 3)

    def test_andina_seleccion_manual_exportaciones_y_cambio_de_excel(self):
        fuente, archivo = self.agregar_andina()
        lista = self.a.listar_tickets(fuente)
        self.assertEqual(len(lista['tickets']), 8)
        self.assertEqual(lista['hoja'], '1.1')
        r = self.solicitud_memoria('POST', '/api/auditoria/seleccionar',
                                  {'fuente': fuente, 'fila': 3, 'sha256': lista['sha256']})
        self.assertEqual(r['status'], 200)
        muestra = json.loads(r['contenido'])
        self.assertEqual(muestra['tickets'][0]['campos'][2]['valor'], '0.2')
        csv_texto = crear_csv(muestra).decode('utf-8-sig')
        self.assertIn('Cantidad (t) [C]', csv_texto)
        for extension in ('pdf', 'csv'):
            self.assertEqual(self.solicitud_memoria('GET', '/api/auditoria/reporte.' + extension + '?id=' + muestra['id'])['status'], 200)
        try:
            from pypdf import PdfReader
        except ImportError:
            PdfReader = None
        if PdfReader:
            pdf = PdfReader(BytesIO(crear_pdf_auditoria(muestra)))
            self.assertIn('0.2 t', '\n'.join(p.extract_text() for p in pdf.pages))
        archivo.write_bytes(excel_andina(cantidad_vacia=True))
        with self.assertRaisesRegex(ValueError, 'Excel cambió'):
            self.a.seleccionar({'fuente': fuente, 'fila': 3, 'sha256': lista['sha256']})

    def test_andina_respaldo_del_servicio_correcto_y_copia_persistente(self):
        fuente, archivo = self.agregar_andina()
        (archivo.parent / '2.5a TICKET INTERNO.pdf').write_bytes(b'%PDF-1.4\nOtro servicio')
        muestra = self.a.sortear({'fuente': fuente, 'cantidad': 3})
        self.assertIsNone(muestra['respaldo_edp'])
        cambio = {'muestra': muestra['id'], 'fila': muestra['tickets'][0]['fila'], 'comprobado': True}
        with self.assertRaisesRegex(ValueError, '1.1 Retiro RINSP'):
            self.a.modificar(cambio, 'comprobar')
        pdf = archivo.parent / '1.1 Retiro RINSP.pdf'
        contenido = b'%PDF-1.4\nRespaldo Andina'
        pdf.write_bytes(contenido)
        muestra = self.a.modificar(cambio, 'comprobar')
        self.assertTrue(muestra['respaldo_edp']['guardado'])
        self.assertTrue(muestra['tickets'][0]['comprobado'])
        pdf.unlink()
        self.assertEqual(self.a.documento_edp(muestra['id'])[1], contenido)

    def solicitud_memoria(self, metodo, ruta, datos=None):
        """Ejecutar el handler sin sockets ni publicar el panel."""
        clase = crear_handler(self.fuentes, self.raiz / 'Datos')
        handler = clase.__new__(clase)
        cuerpo = json.dumps(datos).encode() if datos is not None else b''
        handler.headers = {'Content-Type': 'application/json', 'Content-Length': str(len(cuerpo))}
        handler.rfile = BytesIO(cuerpo)
        resultado = {}
        handler.responder = lambda status, contenido, tipo, descarga=None: resultado.update(
            status=status, contenido=contenido, tipo=tipo, descarga=descarga)
        handler.log_error = lambda *args: None
        if metodo == 'GET':
            handler.auditoria_get(urlsplit(ruta))
        else:
            handler.auditoria_post(ruta)
        return resultado

    def test_seleccion_manual_distingue_repetidos_y_detecta_excel_modificado(self):
        lista = self.a.listar_tickets(self.fuente)
        self.assertEqual(len(lista['tickets']), 8)
        self.assertEqual(lista['tickets'][0]['ticket'], lista['tickets'][1]['ticket'])
        m = self.seleccionar(3)
        self.assertEqual(m['modo'], 'manual')
        self.assertEqual([t['fila'] for t in m['tickets']], [3])
        self.assertEqual(m['tickets'][0]['campos'][3]['valor'], '120')
        for fila in (1, 10000, None, True):
            with self.assertRaises(ValueError):
                self.a.seleccionar({'fuente': self.fuente, 'fila': fila, 'sha256': lista['sha256']})
        self.archivo.write_bytes(excel(9))
        with self.assertRaisesRegex(ValueError, 'Excel cambió'):
            self.a.seleccionar({'fuente': self.fuente, 'fila': 3, 'sha256': lista['sha256']})
        self.assertEqual(self.a.obtener(m['id'])['poblacion'], 8)

    def test_pdf_de_edp_habilita_ambos_modos_y_guarda_una_copia_por_muestra(self):
        pdf = self.archivo.parent / '2.5a TICKET INTERNO.PDF'
        original = b'%PDF-1.4\nRespaldo comun de prueba'
        pdf.write_bytes(original)
        for m in (self.sortear(), self.seleccionar()):
            self.assertEqual(m['respaldo_edp']['nombre'], pdf.name)
            self.assertFalse(any(t['comprobado'] for t in m['tickets']))
            self.assertFalse(any(t['documentos'] for t in m['tickets']))
            for t in m['tickets']:
                m = self.a.modificar({'muestra': m['id'], 'fila': t['fila'], 'comprobado': True}, 'comprobar')
            self.assertTrue(m['respaldo_edp']['guardado'])
            self.assertTrue(all(t['comprobado'] for t in m['tickets']))
            with self.a.conexion() as db:
                self.assertEqual(db.execute('SELECT count(*) FROM documentos WHERE muestra=? AND fila=0', (m['id'],)).fetchone()[0], 1)
        pdf.unlink()
        self.assertEqual(self.a.documento_edp(m['id'])[1], original)
        self.assertTrue(self.a.obtener(m['id'])['tickets'][0]['comprobado'])

    def test_pdf_solo_del_edp_correcto_y_detectado_en_revision_existente(self):
        otro = self.fuentes / 'Codelco El Salvador' / '2026 -06 EDP 48'
        otro.mkdir()
        (otro / '2.5a TICKET INTERNO.pdf').write_bytes(b'%PDF-1.4\notro EDP')
        m = self.seleccionar()
        self.assertIsNone(m['respaldo_edp'])
        cambio = {'muestra': m['id'], 'fila': m['tickets'][0]['fila'], 'comprobado': True}
        with self.assertRaises(ValueError): self.a.modificar(cambio, 'comprobar')
        pdf = self.archivo.parent / '2.5a ticket interno.pdf'
        pdf.write_bytes(b'')
        self.assertIsNone(self.a.obtener(m['id'])['respaldo_edp'])
        pdf.write_bytes(b'%PDF-1.4\nEste EDP')
        self.assertIsNotNone(self.a.obtener(m['id'])['respaldo_edp'])
        respuesta = self.solicitud_memoria('GET', '/api/auditoria/respaldo-edp?id=' + m['id'])
        self.assertEqual(respuesta['status'], 200)
        self.assertEqual(respuesta['contenido'], pdf.read_bytes())
        pdf.unlink()
        with self.assertRaises(ValueError): self.a.modificar(cambio, 'comprobar')

    def test_pdf_con_sufijo_edp_solo_habilita_el_periodo_y_servicio_correcto(self):
        for fuente, archivo in ((self.fuente, self.archivo), self.agregar_andina()):
            m = self.a.sortear({'fuente': fuente, 'cantidad': 3})
            base = '1.1 Retiro RINSP' if m['unidad'] == 'Andina' else '2.5a TICKET INTERNO'
            equivocado = archivo.parent / f'{base} (EDP {m["edp"] + 1}).pdf'
            equivocado.write_bytes(b'%PDF-1.4\nOtro EDP')
            self.assertIsNone(self.a.obtener(m['id'])['respaldo_edp'])
            correcto = archivo.parent / f'{base} ( EDP {m["edp"]}).PDF'
            correcto.write_bytes(b'%PDF-1.4\nEste EDP')
            cambio = {'muestra': m['id'], 'fila': m['tickets'][0]['fila'], 'comprobado': True}
            m = self.a.modificar(cambio, 'comprobar')
            self.assertTrue(m['tickets'][0]['comprobado'])
            self.assertFalse(m['tickets'][0]['documentos'])
            self.assertEqual(m['respaldo_edp']['nombre'], correcto.name)
            correcto.unlink()
            self.assertEqual(self.a.documento_edp(m['id'])[1], b'%PDF-1.4\nEste EDP')

    def test_cobertura_no_duplica_edp_y_actualiza_chequeo(self):
        fuente_andina, _ = self.agregar_andina()
        otro = self.fuentes / 'Codelco El Salvador' / '2026 -06 EDP 48'
        otro.mkdir()
        (otro / 'Otro.xlsx').write_bytes(excel())
        # Una segunda versión de Excel no debe duplicar el EDP en el gráfico.
        (self.archivo.parent / 'Version.xlsx').write_bytes(excel())
        (self.archivo.parent / '2.5a TICKET INTERNO (EDP 47).pdf').write_bytes(b'%PDF-1.4\nEste EDP')
        m = self.seleccionar()
        segunda = self.seleccionar(4)
        m_andina = self.a.sortear({'fuente': fuente_andina, 'cantidad': 3})
        cambio = {'muestra': m['id'], 'fila': m['tickets'][0]['fila'], 'comprobado': True}
        self.a.modificar(cambio, 'comprobar')
        registros = self.a.catalogo()['estados_pago']
        self.assertEqual(len(registros), 3)
        edp = next(e for e in registros if e['unidad'] == 'El Salvador' and e['edp'] == 47)
        self.assertEqual((edp['revisiones'], edp['comprobados'], edp['pendientes']), (2, 0, 1))
        self.assertEqual(edp['muestra_vigente'], segunda['id'])
        self.assertEqual((edp['estado_codigo'], edp['color']), ('sin_revision', 'rojo'))
        self.assertFalse(edp['con_chequeo'])
        self.assertEqual(len(edp['archivos']), 2)
        self.a.modificar({**cambio, 'comprobado': False}, 'comprobar')
        edp = next(e for e in self.a.catalogo()['estados_pago'] if e['unidad'] == 'El Salvador' and e['edp'] == 47)
        self.assertFalse(edp['con_chequeo'])
        self.assertEqual(edp['estado'], 'Sin revisión')
        self.assertEqual((edp['estado_codigo'], edp['color']), ('sin_revision', 'rojo'))
        self.a.modificar(cambio, 'comprobar')
        self.a.modificar({'muestra': segunda['id'], 'fila': segunda['tickets'][0]['fila'], 'comprobado': True}, 'comprobar')
        reporte = self.a.reporte_unidad('El Salvador')
        self.assertEqual(len(reporte['estados_pago']), 2)
        self.assertEqual(len(reporte['muestras']), 2)
        self.assertTrue(all(x['unidad'] == 'El Salvador' for x in reporte['muestras']))
        self.assertEqual(reporte['estados_pago'][0]['estado'], 'Chequeo completo')
        self.assertEqual(reporte['estados_pago'][0]['color'], 'verde')
        self.assertEqual(reporte['estados_pago'][1]['estado'], 'Sin revisión')
        self.archivo.unlink()
        (self.archivo.parent / 'Version.xlsx').unlink()
        self.assertEqual(len(self.a.reporte_unidad('El Salvador')['estados_pago']), 2)
        with self.assertRaises(ValueError): self.a.reporte_unidad('desconocida')
        self.assertNotIn(m_andina['id'], [x['id'] for x in reporte['muestras']])

    def test_revision_nueva_completa_no_suma_pendientes_del_historial(self):
        (self.archivo.parent / '2.5a TICKET INTERNO.pdf').write_bytes(b'%PDF-1.4\nEste EDP')
        antigua = self.sortear()
        vigente = self.sortear()
        for ticket in vigente['tickets']:
            self.a.modificar({'muestra': vigente['id'], 'fila': ticket['fila'], 'comprobado': True}, 'comprobar')
        edp = self.a.catalogo()['estados_pago'][0]
        self.assertEqual((edp['tickets'], edp['comprobados'], edp['pendientes']), (3, 3, 0))
        self.assertEqual((edp['estado_codigo'], edp['color']), ('completo', 'verde'))
        self.assertEqual(edp['muestra_vigente'], vigente['id'])
        self.assertEqual(edp['revisiones'], 2)
        reporte = self.a.reporte_unidad('El Salvador')
        self.assertEqual(reporte['estados_pago'][0], edp)
        self.assertEqual({m['id'] for m in reporte['muestras']}, {antigua['id'], vigente['id']})
        self.assertFalse(any(t['comprobado'] for t in self.a.obtener(antigua['id'])['tickets']))
        # Editar el historial más tarde no sustituye la revisión vigente.
        self.a.modificar({'muestra': antigua['id'], 'fila': antigua['tickets'][0]['fila'], 'comprobado': True}, 'comprobar')
        self.assertEqual(self.a.catalogo()['estados_pago'][0], edp)
        cambio = {'muestra': vigente['id'], 'fila': vigente['tickets'][0]['fila'], 'comprobado': False}
        self.a.modificar(cambio, 'comprobar')
        edp = self.a.catalogo()['estados_pago'][0]
        self.assertEqual((edp['comprobados'], edp['pendientes'], edp['color']), (2, 1, 'amarillo'))
        nueva = self.sortear()
        edp = self.a.catalogo()['estados_pago'][0]
        self.assertEqual((edp['comprobados'], edp['pendientes'], edp['color']), (0, 3, 'rojo'))
        self.assertEqual(edp['muestra_vigente'], nueva['id'])
        self.assertEqual(self.a.catalogo()['muestras'][0]['id'], nueva['id'])

    def test_revision_vigente_independiente_del_orden_de_sincronizacion(self):
        base = dict(unidad='El Salvador', periodo='2026-08', anio=2026, mes=8, edp=50,
                    archivo='Agosto.xlsx', cantidad=3)
        vieja = {**base, 'id': 'vieja', 'creado': '2026-10-05T10:59:59-03:00', 'comprobados': 0}
        nueva = {**base, 'id': 'nueva', 'creado': '2026-10-05T14:00:00+00:00', 'comprobados': 3}
        primero = resumen_edp([], [vieja, nueva])
        self.assertEqual(primero, resumen_edp([], [nueva, vieja]))
        self.assertEqual((primero[0]['muestra_vigente'], primero[0]['color']), ('nueva', 'verde'))
        empate = {**nueva, 'id': 'z', 'comprobados': 1}
        self.assertEqual(resumen_edp([], [nueva, empate]), resumen_edp([], [empate, nueva]))
        self.assertEqual(resumen_edp([], [empate, nueva])[0]['color'], 'amarillo')

    def test_pdf_unidad_incluye_periodos_sin_muestra_y_todos_los_campos(self):
        self.agregar_andina()
        otro = self.fuentes / 'Codelco El Salvador' / '2025 -12 EDP 42'
        otro.mkdir()
        (otro / 'Sin_revision.xlsx').write_bytes(excel())
        (self.archivo.parent / '2.5a TICKET INTERNO (EDP 47).pdf').write_bytes(b'%PDF-1.4\nEste EDP')
        m = self.seleccionar()
        self.a.modificar({'muestra': m['id'], 'fila': m['tickets'][0]['fila'], 'comprobado': True}, 'comprobar')
        contenido = crear_pdf_unidad(self.a.reporte_unidad('El Salvador'))
        self.assertTrue(contenido.startswith(b'%PDF'))
        r = self.solicitud_memoria('GET', '/api/auditoria/unidad.pdf?unidad=El%20Salvador')
        self.assertEqual((r['status'], r['tipo']), (200, 'application/pdf'))
        self.assertEqual(r['descarga'], 'el-salvador-varios-meses-2025-12-a-2026-05-edp-42-47-panel-auditoria-unidad.pdf')
        for unidad in ('', 'Otro'):
            self.assertEqual(self.solicitud_memoria('GET', '/api/auditoria/unidad.pdf?unidad=' + unidad)['status'], 400)
        try:
            from pypdf import PdfReader
        except ImportError:
            self.skipTest('pypdf opcional no disponible.')
        textos = '\n'.join(p.extract_text() for p in PdfReader(BytesIO(contenido)).pages)
        for esperado in ('2025-12', '2026-05', 'Sin revisiones guardadas', m['id'], m['sha256'],
                         'Campo adicional', 'COMPROBADO MANUALMENTE', '2.5a TICKET INTERNO (EDP 47).pdf'):
            self.assertIn(esperado, textos)
        self.assertNotIn('Andina', textos)

    def test_adjuntos_30mb_en_ambos_modos_y_rechazo_superior(self):
        archivo = adjunto('respaldo.pdf', b'x' * MAX_ARCHIVO)
        for m in (self.sortear(), self.seleccionar()):
            datos = {**archivo, 'muestra': m['id'], 'fila': m['tickets'][0]['fila']}
            self.assertLess(len(json.dumps(datos).encode()), MAX_SOLICITUD)
            respuesta = self.solicitud_memoria('POST', '/api/auditoria/adjuntar', datos)
            self.assertEqual(respuesta['status'], 200)
            guardado = json.loads(respuesta['contenido'])
            self.assertEqual(guardado['tickets'][0]['documentos'][0]['bytes'], MAX_ARCHIVO)
        with self.assertRaisesRegex(ValueError, '30 MB'):
            decodificar(adjunto('grande.pdf', b'x' * (MAX_ARCHIVO + 1)), EXT_DOCUMENTOS)

    def test_rutas_seleccion_y_carga_excel_retirada_sin_servidor(self):
        r = self.solicitud_memoria('GET', '/api/auditoria/tickets?id=' + self.fuente)
        self.assertEqual(r['status'], 200)
        lista = json.loads(r['contenido'])
        r = self.solicitud_memoria('POST', '/api/auditoria/seleccionar',
                                   {'fuente': self.fuente, 'fila': 3, 'sha256': lista['sha256']})
        self.assertEqual(r['status'], 200)
        m = json.loads(r['contenido'])
        self.assertEqual(m['modo'], 'manual')
        for extension in ('pdf', 'csv'):
            self.assertEqual(self.solicitud_memoria('GET', '/api/auditoria/reporte.' + extension + '?id=' + m['id'])['status'], 200)
        self.assertEqual(self.solicitud_memoria('POST', '/api/auditoria/importar', {})['status'], 404)

    def test_exportaciones_manual_respaldo_comun_sin_formula_peso(self):
        (self.archivo.parent / '2.5a TICKET INTERNO.pdf').write_bytes(b'%PDF-1.4\nPrueba')
        m = self.seleccionar()
        m['tickets'][0]['campos'][3]['formula'] = '=Tabla642[[#This Row],[Peso en Báscula (Kg)]]-Tabla642[[#This Row],[Tara Camión (Kg)]]'
        contenido = crear_csv(m).decode('utf-8-sig')
        self.assertIn('2.5a TICKET INTERNO.pdf', contenido)
        self.assertIn('manual', contenido)
        self.assertNotIn('Tabla642', contenido)
        try:
            from pypdf import PdfReader
        except ImportError:
            self.skipTest('pypdf opcional no disponible para extraer el texto del PDF.')
        pdf = PdfReader(BytesIO(crear_pdf_auditoria(m)))
        contenido = '\n'.join(p.extract_text() for p in pdf.pages)
        self.assertIn('Manual: 1', contenido)
        self.assertIn('2.5a TICKET INTERNO.pdf', contenido)
        self.assertIn('120', contenido)
        self.assertNotIn('Tabla642', contenido)

    def test_lectura_todas_las_columnas_ocultas_repetidos_y_formulas(self):
        for tabla in (True, False):
            datos = leer_excel(excel(tabla=tabla, hoja='RES.NO PELIGROSOS'))
            self.assertEqual(len(datos['tickets']), 8)
            self.assertEqual(len(datos['tickets'][0]['campos']), 6)
            self.assertEqual(datos['tickets'][2]['ticket'], '000002')
            self.assertEqual(datos['tickets'][0]['ticket'], '0')
            self.assertEqual(datos['tickets'][1]['fila'], 3)
            self.assertTrue(datos['tickets'][0]['campos'][4]['sin_resultado'])
            self.assertTrue(any('repetidos' in a for a in datos['avisos']))

    def test_muestras_3_4_5_sin_repetir_filas_y_sin_imputaciones(self):
        for cantidad in (3, 4, 5):
            m = self.sortear(cantidad)
            self.assertEqual(len(m['tickets']), cantidad)
            self.assertEqual(len({t['fila'] for t in m['tickets']}), cantidad)
            self.assertEqual(m['periodo'], '2026-05')  # Fechas abril pertenecen al EDP mayo.
            self.assertEqual(m['poblacion'], 8)
        self.assertEqual(len(self.a.catalogo()['muestras']), 3)

    def test_tamano_invalido_y_poblacion_insuficiente_no_guarda(self):
        for cantidad in (2, 6, True, '3', None):
            with self.assertRaises(ValueError): self.sortear(cantidad)
        self.archivo.write_bytes(excel(2))
        with self.assertRaisesRegex(ValueError, 'no alcanza'): self.sortear(3)
        self.assertEqual(self.a.catalogo()['muestras'], [])

    def test_documentos_persisten_y_comprobacion_solo_manual(self):
        m = self.sortear()
        cambio = {'muestra': m['id'], 'fila': m['tickets'][0]['fila']}
        with self.assertRaisesRegex(ValueError, 'Adjunta'):
            self.a.modificar({**cambio, 'comprobado': True}, 'comprobar')
        m = self.a.modificar({**cambio, **adjunto()}, 'adjuntar')
        self.assertFalse(m['tickets'][0]['comprobado'])
        m = self.a.modificar({**cambio, 'comprobado': True}, 'comprobar')
        self.assertTrue(m['tickets'][0]['comprobado'])
        reinicio = Auditoria(self.fuentes, self.raiz / 'Datos')
        self.assertEqual(reinicio.obtener(m['id']), m)
        doc = m['tickets'][0]['documentos'][0]
        self.assertEqual(reinicio.documento(doc['id'])[1], b'%PDF-1.4\nRespaldo de prueba')
        # Cambiar la fuente no altera lo que se revisó.
        self.archivo.write_bytes(excel(2))
        self.assertEqual(reinicio.obtener(m['id'])['poblacion'], 8)
        m = self.a.modificar({**cambio, **adjunto('segundo.txt', b'Otro respaldo')}, 'adjuntar')
        self.assertFalse(m['tickets'][0]['comprobado'])
        self.a.modificar({**cambio, 'comprobado': True}, 'comprobar')
        m = self.a.modificar({**cambio, 'documento': doc['id']}, 'quitar')
        self.assertFalse(m['tickets'][0]['comprobado'])

    def test_aislamiento_entre_tickets_y_muestras(self):
        m, otro = self.sortear(), self.sortear()
        cambio = {'muestra': m['id'], 'fila': m['tickets'][0]['fila']}
        m = self.a.modificar({**cambio, **adjunto()}, 'adjuntar')
        doc = m['tickets'][0]['documentos'][0]
        with self.assertRaises(ValueError):
            self.a.modificar({'muestra': otro['id'], 'fila': otro['tickets'][0]['fila'], 'documento': doc['id']}, 'quitar')
        self.assertFalse(any(t['documentos'] for t in self.a.obtener(otro['id'])['tickets']))
        self.assertFalse(any(t['documentos'] for t in m['tickets'][1:]))

    def test_cambios_concurrentes_no_pierden_documentos(self):
        m = self.sortear()
        def subir(ticket):
            self.a.modificar({'muestra': m['id'], 'fila': ticket['fila'], **adjunto()}, 'adjuntar')
        with ThreadPoolExecutor(max_workers=3) as executor:
            list(executor.map(subir, m['tickets']))
        self.assertTrue(all(len(t['documentos']) == 1 for t in self.a.obtener(m['id'])['tickets']))

    def test_importacion_contexto_y_archivos_invalidos(self):
        datos = {**adjunto('EDP.xlsx', excel_andina()), 'unidad': 'Andina', 'anio': 2025, 'mes': 8, 'edp': 12}
        fuente = self.a.importar(datos)
        self.assertEqual(self.a.importar(datos)['id'], fuente['id'])
        self.assertEqual(len(self.a.catalogo()['fuentes']), 2)
        m = self.a.sortear({'fuente': fuente['id'], 'cantidad': 5})
        self.assertEqual((m['unidad'], m['periodo'], m['edp']), ('Andina', '2025-08', 12))
        for parche in ({'mes': 13}, {'unidad': ''}, {'nombre': '../EDP.xlsx'}, {'contenido': 'mal-base64'}, {'edp': '12'}):
            with self.assertRaises(ValueError): self.a.importar({**datos, **parche})
        with self.assertRaises(ValueError): self.a.importar({**datos, **adjunto('EDP.xlsx', b'no excel')})

    def test_csv_y_pdf_tienen_muestra_estados_y_todas_las_columnas(self):
        m = self.sortear(5)
        m['tickets'][0]['campos'][-1]['valor'] = '=HYPERLINK("https://example.com")'
        contenido = crear_csv(m)
        self.assertTrue(contenido.startswith(b'\xef\xbb\xbf'))
        filas = list(csv.reader(StringIO(contenido.decode('utf-8-sig')), delimiter=';'))
        self.assertEqual(len(filas), 6)
        self.assertIn('Campo adicional [F]', filas[0])
        self.assertTrue(filas[1][filas[0].index('Campo adicional [F]')].startswith("'="))
        self.assertEqual(filas[1][filas[0].index('Estado revisión manual')], 'Pendiente de comprobación')
        pdf = crear_pdf_auditoria(m)
        self.assertTrue(pdf.startswith(b'%PDF'))
        try:
            from pypdf import PdfReader
        except ImportError:
            return
        lector = PdfReader(BytesIO(pdf))
        self.assertEqual(len(lector.pages), 6)
        textos = '\n'.join(p.extract_text() for p in lector.pages)
        self.assertIn(m['id'], textos)
        self.assertIn('Campo adicional', textos)
        self.assertIn('PENDIENTE DE COMPROBACIÓN', textos)

    def test_http_flujo_completo_y_rutas_privadas(self):
        servidor = ServidorPanel(('127.0.0.1', 0), crear_handler(self.fuentes, self.raiz / 'Datos'))
        hilo = threading.Thread(target=servidor.serve_forever, daemon=True); hilo.start()
        base = f'http://127.0.0.1:{servidor.server_port}'
        def post(ruta, datos, origen=base):
            with urlopen(Request(base + '/api/auditoria/' + ruta, json.dumps(datos).encode(),
                                 {'Content-Type': 'application/json', 'Origin': origen})) as r:
                return json.load(r)
        try:
            m = post('sortear', {'fuente': self.fuente, 'cantidad': 3})
            cambio = {'muestra': m['id'], 'fila': m['tickets'][0]['fila']}
            post('adjuntar', {**cambio, **adjunto('guía.pdf')})
            m = post('comprobar', {**cambio, 'comprobado': True})
            for ruta, tipo in [('reporte.pdf', 'application/pdf'), ('reporte.csv', 'text/csv')]:
                with urlopen(base + '/api/auditoria/' + ruta + '?id=' + m['id']) as r:
                    self.assertIn(tipo, r.headers['Content-Type']); self.assertTrue(r.read())
            doc = m['tickets'][0]['documentos'][0]
            with urlopen(base + '/api/auditoria/documento?id=' + doc['id']) as r:
                self.assertIn('filename*=', r.headers['Content-Disposition'])
            for ruta in ('/Datos/Auditoria/auditoria.sqlite3', '/Fuentes/EDP.xlsx'):
                with self.assertRaises(HTTPError) as error: urlopen(base + ruta)
                self.assertEqual(error.exception.code, 404)
            with self.assertRaises(HTTPError) as error:
                # El origen se rechaza antes de leer el cuerpo. Sin bytes pendientes,
                # Windows no convierte el cierre del socket en un reset intermitente.
                urlopen(Request(base + '/api/auditoria/sortear', b'',
                                {'Content-Type': 'application/json', 'Origin': 'https://otro.example'}))
            self.assertEqual(error.exception.code, 400)
        finally:
            servidor.shutdown(); servidor.server_close(); hilo.join()


if __name__ == '__main__':
    unittest.main()
